from __future__ import annotations

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import APIRouter, FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_classification_result_openapi import _bind_result_standard
from test_classification_result_pool import _publish, _seed_result_context
from test_manual_correction import _items, _record
from test_result_version_reviews import _publish_review_required

from web_backend.classification_result_service import ClassificationResultService
from web_backend.classification_results.manual_correction import correct_group
from web_backend.classification_results.result_export import (
    SEMANTIC_COLUMNS,
    SOURCE_COLUMNS,
    SUMMARY_COLUMNS,
    semantic_export_rows,
)
from web_backend.classification_standard_service import ClassificationStandardService
from web_backend.routers.task_download_routes import register_task_download_routes
from web_backend.task_service import TaskService

SHEETS = ["反馈汇总", "语义明细", "源明细"]


@pytest.fixture
def context(tmp_path: Path) -> SimpleNamespace:
    value = _seed_result_context(tmp_path)
    value.version = _publish(value)
    _bind_result_standard(
        ClassificationStandardService(value.database),
        value.version["result_id"],
        value.taxonomy.version,
    )
    value.service = ClassificationResultService(value.database)
    return value


def _rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row, strict=True)) for row in values[1:]]


def test_workbook_groups_match_page_and_preserve_all_sources(context):
    with context.database.transaction() as connection:
        connection.execute(
            "UPDATE classification_result_records SET fnsku='FIRST-FNSKU' WHERE source_row=2"
        )
        connection.execute(
            "UPDATE classification_result_records SET fnsku='SECOND-FNSKU' WHERE source_row=3"
        )
    version_id = context.version["version_id"]
    content, filename = context.service.download(version_id)
    book = load_workbook(BytesIO(content))
    assert filename == "classification-L1-v1.xlsx"
    assert book.sheetnames == SHEETS
    summaries, semantics, sources = [_rows(book[name]) for name in SHEETS]
    groups = context.service.record_groups(version_id)["items"]
    assert len(summaries) == len(groups) == 2
    assert len(sources) == 3
    assert [row["源文件行号"] for row in sources] == [2, 3, 4]
    for row, group in zip(summaries, groups, strict=True):
        record = group["record"]
        assert row["源明细数"] == group["member_count"]
        assert row["产品名称"] == record["product_name"]
        assert row["反馈正文"] == record["comment"]
        assert row["结果状态"] == "可用"
        assert row["是否需重跑"] == "否"
        details = [item for item in semantics if item["反馈编号"] == row["反馈编号"]]
        assert len(details) == len(
            record["classification"]["semantic_review"]["semantic_items"]
        )
        assert row["中文语义"] == "\n".join(
            f"{item['语义序号']}. {item['中文语义']}" for item in details
        )
        assert row["分类路径"] == "\n".join(
            f"{item['语义序号']}. {item['分类路径']}" for item in details
        )
        assert (
            len([s for s in sources if s["反馈编号"] == row["反馈编号"]])
            == group["member_count"]
        )
    assert sources[0]["反馈编号"] == sources[1]["反馈编号"]
    assert [row["FNSKU"] for row in sources[:2]] == ["FIRST-FNSKU", "SECOND-FNSKU"]
    for sheet, columns in zip(
        book, [SUMMARY_COLUMNS, SEMANTIC_COLUMNS, SOURCE_COLUMNS], strict=True
    ):
        assert list(next(sheet.values)) == columns
        assert sheet.freeze_panes == ("D2" if sheet.title == "反馈汇总" else "C2")
        assert sheet.auto_filter.ref == sheet.dimensions
        assert not sheet.merged_cells
        assert sheet["A2"].alignment.wrap_text
        assert sheet.column_dimensions["F"].width >= 20
    assert all("classification_json" not in next(sheet.values) for sheet in book)
    book.close()


def test_empty_workbook_keeps_headers_and_fallback_filename(context):
    with context.database.transaction() as connection:
        connection.execute("DELETE FROM classification_result_records")
        connection.execute("UPDATE classification_results SET listing = NULL")
    content, filename = context.service.download(context.version["version_id"])
    assert filename == f"classification-{context.version['result_id']}-v1.xlsx"
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == SHEETS
    assert all(sheet.max_row == 1 and not _rows(sheet) for sheet in book)
    book.close()


def test_semantic_rows_keep_multiple_opinions_evidence_and_review_reasons():
    item = {
        "classification": {
            "semantic_review": {
                "semantic_items": [
                    {
                        "opinion": "尺码偏小",
                        "label_code": "A",
                        "label_path": ["尺码", "偏小"],
                        "evidence_text": "too small",
                        "disposition": "MAPPED",
                    },
                    {
                        "opinion": "另一条待判断语义",
                        "evidence_text": "原文片段",
                        "disposition": "TRUE_AMBIGUITY",
                        "reason": "含义存在歧义",
                    },
                ]
            }
        }
    }
    rows = semantic_export_rows(item, "F000001")
    assert [row["语义序号"] for row in rows] == [1, 2]
    assert rows[0]["分类路径"] == "尺码 → 偏小"
    assert rows[0]["原文证据"] == "too small"
    assert rows[1]["处置状态"] == "原文含义不明确"
    assert rows[1]["判断说明"] == "含义存在歧义"
    assert semantic_export_rows({"classification": {}}, "F000002") == []
    unexplained = semantic_export_rows(
        {
            "classification": {
                "semantic_review": {"unexplained_fragments": ["剩余原文"]}
            }
        },
        "F000003",
    )
    assert unexplained[0]["原文证据"] == "剩余原文"
    assert unexplained[0]["处置状态"] == "未解释原文片段"


def test_download_keeps_three_reads_without_member_queries(context, monkeypatch):
    statements = []
    connect = context.database.connect

    def traced_connection():
        connection = connect()
        connection.set_trace_callback(statements.append)
        return connection

    monkeypatch.setattr(context.database, "connect", traced_connection)
    content, _ = context.service.download(context.version["version_id"])
    assert content.startswith(b"PK")
    assert (
        len([sql for sql in statements if sql.lstrip().upper().startswith("SELECT")])
        == 3
    )


def test_task_downloads_use_current_manual_result_instead_of_checkpoint(tmp_path):
    context, version = _publish_review_required(tmp_path)
    service = ClassificationResultService(context.database)
    correct_group(
        service, version["version_id"], _record(context, version), _items(), "user-2"
    )
    tasks = TaskService(context.database)
    segment = tasks.get(context.task_id)["segments"][0]
    router = APIRouter()
    register_task_download_routes(router, tasks, lambda: {"id": "user-1"})
    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)
    for path in [
        f"/api/tasks/{context.task_id}/download",
        f"/api/tasks/{context.task_id}/segments/{segment['segment_key']}/download",
    ]:
        response = client.get(path)
        assert response.status_code == 200
        book = load_workbook(BytesIO(response.content))
        for record in _rows(book["反馈汇总"]):
            if record["订单号"] == "ORDER-DUP":
                assert "整体尺寸偏小" in record["中文语义"]
                assert record["结果状态"] == "可用"
                assert record["人工修正情况"] == "人工已修正"
                assert record["修正人"]
                assert record["是否需重跑"] == "否"
            else:
                assert record["结果状态"] == "需复核"
        book.close()


def test_combined_download_has_unique_feedback_numbers(context):
    content = context.service.download_versions([context.version["version_id"]] * 2)
    book = load_workbook(BytesIO(content))
    summaries = _rows(book["反馈汇总"])
    assert len(summaries) == len({row["反馈编号"] for row in summaries}) == 4
    assert len(_rows(book["源明细"])) == 6
    book.close()


@pytest.mark.parametrize(
    "quality, rerun",
    [("ready", 0), ("review_required", 0), ("unusable", 1), ("excluded", 0)],
)
def test_feedback_without_semantics_or_with_failure_is_not_lost(
    context, quality, rerun
):
    with context.database.transaction() as connection:
        connection.execute(
            "UPDATE classification_units SET classification_json='{}', processing_status=?, system_rerun_required=?",
            ("MODEL_ERROR" if rerun else "AUTO_APPROVED", rerun),
        )
        connection.execute(
            "UPDATE classification_result_records SET quality_status=?", (quality,)
        )
    content, _ = context.service.download(context.version["version_id"])
    book = load_workbook(BytesIO(content))
    assert len(_rows(book["反馈汇总"])) == 2
    assert len(_rows(book["源明细"])) == 3
    expected = {
        "ready": "可用",
        "review_required": "需复核",
        "unusable": "不可用",
        "excluded": "已忽略",
    }
    assert all(
        row["结果状态"] == expected[quality]
        and row["是否需重跑"] == ("是" if rerun else "否")
        for row in _rows(book["反馈汇总"])
    )
    if rerun:
        assert "系统处理失败" in {row["处置状态"] for row in _rows(book["语义明细"])}
    book.close()


def test_blank_orders_stay_separate_and_long_text_is_literal(context):
    text = "=原文不是公式\n" + "完整反馈 " * 600
    with context.database.transaction() as connection:
        connection.execute(
            "UPDATE classification_result_records SET order_id='', comment=?", (text,)
        )
    content, _ = context.service.download(context.version["version_id"])
    book = load_workbook(BytesIO(content))
    assert len(_rows(book["反馈汇总"])) == 3
    for title in ["反馈汇总", "源明细"]:
        sheet = book[title]
        column = list(next(sheet.values)).index("反馈正文") + 1
        assert all(
            sheet.cell(index, column).value == text
            and sheet.cell(index, column).data_type == "s"
            for index in range(2, 5)
        )
    book.close()
