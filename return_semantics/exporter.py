from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from return_semantics.data import ReturnDataset
from return_semantics.export_detail import _build_detail_rows
from return_semantics.export_formatting import _display_key
from return_semantics.export_review import _build_semantic_review_rows
from return_semantics.export_semantics import (
    _build_dimension_decision_rows,
    _build_semantic_rows,
    _build_statistics,
    _build_unknown_rows,
)
from return_semantics.schemas import (
    ProcessingStatus,
    TaxonomyConfig,
    ValidatedClassification,
)
from return_semantics.semantic_review import (
    BUSINESS_REVIEW_REQUIRED,
    READY,
    SYSTEM_RERUN_REQUIRED,
    review_route,
)

REVIEW_STATUSES = {
    ProcessingStatus.SECONDARY_REVIEW.value,
    ProcessingStatus.MANUAL_REVIEW.value,
    ProcessingStatus.UNKNOWN_SEMANTIC.value,
    ProcessingStatus.MODEL_ERROR.value,
}


def _style_workbook(writer: pd.ExcelWriter) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in writer.book.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_index, column_cells in enumerate(sheet.columns, 1):
            values = [str(cell.value or "") for cell in column_cells[:200]]
            width = min(max(max(map(len, values), default=10) + 2, 12), 50)
            sheet.column_dimensions[get_column_letter(column_index)].width = width


def build_analysis_frames(
    dataset: ReturnDataset,
    results: dict[str, ValidatedClassification],
    taxonomy: TaxonomyConfig,
) -> tuple[dict[str, pd.DataFrame], dict[str, str]]:
    detail = pd.DataFrame(_build_detail_rows(dataset, results, taxonomy))
    semantics = pd.DataFrame(_build_semantic_rows(dataset, results, taxonomy))
    unknown = pd.DataFrame(_build_unknown_rows(dataset, results))
    source_comments = dataset.unique_comments.set_index("classification_key")[
        "comment_normalized"
    ]
    routes = {
        _display_key(key): review_route(
            result, str(source_comments.loc[key] or ""), taxonomy
        )
        for key, result in results.items()
    }
    detail["复核路由"] = (
        detail["分类键"]
        .map(routes)
        .mask(
            detail["处理状态"].isin({"EXCLUDED_MISSING_CATEGORY", "NO_TEXT_EVIDENCE"}),
            READY,
        )
    )
    return {"分类明细": detail, "语义单元": semantics, "未知语义": unknown}, routes


def export_results(
    output_path: Path,
    dataset: ReturnDataset,
    results: dict[str, ValidatedClassification],
    taxonomy: TaxonomyConfig,
) -> None:
    frames, routes = build_analysis_frames(dataset, results, taxonomy)
    detail, semantics, unknown = (
        frames[name] for name in ("分类明细", "语义单元", "未知语义")
    )
    semantic_review = pd.DataFrame(
        _build_semantic_review_rows(dataset, results, taxonomy)
    )
    if semantic_review.empty:
        business_review = semantic_review.copy()
        system_rerun = semantic_review.copy()
    else:
        # 以评论的最终处置分流，保留其完整语义上下文供阅读核对。
        review_routes = semantic_review["分类键"].map(routes)
        system_rerun = semantic_review.loc[
            review_routes.eq(SYSTEM_RERUN_REQUIRED)
        ].copy()
        business_review = semantic_review.loc[
            review_routes.eq(BUSINESS_REVIEW_REQUIRED)
        ].copy()
    decisions = pd.DataFrame(_build_dimension_decision_rows(results))
    statistics = pd.DataFrame(_build_statistics(dataset, results, taxonomy))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        detail.to_excel(writer, sheet_name="分类明细", index=False)
        semantics.to_excel(writer, sheet_name="语义单元", index=False)
        business_review.to_excel(writer, sheet_name="人工复核", index=False)
        system_rerun.to_excel(writer, sheet_name="系统待重跑", index=False)
        unknown.to_excel(writer, sheet_name="未知语义", index=False)
        semantic_review.to_excel(writer, sheet_name="语义核验", index=False)
        decisions.to_excel(writer, sheet_name="维度裁决", index=False)
        statistics.to_excel(writer, sheet_name="标签统计", index=False)
        _style_workbook(writer)
