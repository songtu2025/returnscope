"""为结果下载设置可筛选、可追溯的阅读布局。"""

from __future__ import annotations

import math
from collections.abc import Iterator
from functools import lru_cache
from io import BytesIO
from typing import Any
from unicodedata import east_asian_width

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter

from web_backend.classification_results.result_export import (
    SEMANTIC_COLUMNS,
    SOURCE_COLUMNS,
    SUMMARY_COLUMNS,
)

_WIDTHS = {
    "反馈编号": 16,
    "订单号": 24,
    "产品名称": 28,
    "反馈正文": 54,
    "反馈标题 / 原因": 28,
    "中文语义": 42,
    "分类路径": 46,
    "原文证据": 48,
    "判断说明": 42,
    "修正时间": 27,
    "源记录编号": 36,
}


def result_workbook(
    summaries: list[dict[str, Any]],
    semantics: list[dict[str, Any]],
    sources: list[dict[str, Any]],
) -> bytes:
    workbook = Workbook(write_only=True)
    font = Font(name="微软雅黑", size=11, color="153D34")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    alignment = Alignment(vertical="top", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="176C57")
    stripe_fill = PatternFill("solid", fgColor="F0F7F4")
    styles: dict[tuple[int, tuple[int, ...]], StyleArray] = {}
    display_lines = lru_cache(maxsize=4096)(_display_lines)
    for title, columns, rows in (
        ("反馈汇总", SUMMARY_COLUMNS, summaries),
        ("语义明细", SEMANTIC_COLUMNS, semantics),
        ("源明细", SOURCE_COLUMNS, sources),
    ):
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = "D2" if title == "反馈汇总" else "C2"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        sheet.row_dimensions[1].height = 32
        for index, name in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = _WIDTHS.get(
                name, 20
            )
        for row_index, values in enumerate(_sheet_values(columns, rows), 1):
            cells = []
            kind = 0 if row_index == 1 else (1 if row_index % 2 == 0 else 2)
            for value in values:
                cell: Any = WriteOnlyCell(sheet, value=value)
                # 所有原文和标识均按文本处理，不能被 Excel 当作公式。
                if isinstance(value, str):
                    cell.data_type = "s"
                # 保留值推导出的日期等格式；样式建成后只共享，不再修改。
                initial_style = cell._style
                key = (kind, tuple(initial_style) if initial_style else ())
                style = styles.get(key)
                if style is None:
                    cell.font = header_font if kind == 0 else font
                    cell.alignment = alignment
                    if kind < 2:
                        cell.fill = header_fill if kind == 0 else stripe_fill
                    styles[key] = cell._style
                else:
                    cell._style = style
                cells.append(cell)
            if row_index > 1:
                lines = max(
                    display_lines(str(value or ""), _WIDTHS.get(name, 20))
                    for name, value in zip(columns, values, strict=True)
                )
                sheet.row_dimensions[row_index].height = min(
                    max(lines * 16 + 10, 42), 180
                )
            sheet.append(cells)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _display_lines(value: Any, width: int) -> int:
    return sum(
        max(
            1,
            math.ceil(
                sum(2 if east_asian_width(c) in "WF" else 1 for c in line) / width
            ),
        )
        for line in str(value or "").split("\n")
    )


def _sheet_values(
    columns: list[str], rows: list[dict[str, Any]]
) -> Iterator[list[Any]]:
    yield list(columns)
    for row in rows:
        yield [row.get(name) for name in columns]
