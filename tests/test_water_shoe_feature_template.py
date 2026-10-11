from __future__ import annotations

import csv
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.generate_water_shoe_feature_template import (
    aggregate_skus,
    build_asin_sku_map,
)
from scripts.water_shoe_workbook import build_workbook, verify_workbook


def _write_source(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", encoding="cp1252", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["return-date", "sku", "asin", "fnsku", "product-name"])
        writer.writerows(rows)


@pytest.fixture
def source(tmp_path: Path) -> Path:
    path = tmp_path / "synthetic.csv"
    _write_source(
        path,
        [
            ["2026-01-01", "SKU-Z", "ASIN-Z", "FNSKU-Z", "Old title"],
            ["2026-02-01", "SKU-Z", "ASIN-Z", "FNSKU-Z", "Café latest"],
            ["2026-03-01", "Amazon.Found.ASIN-Z", "ASIN-Z", "ASIN-Z", "Alias"],
            ["2026-01-01", "sku-a", "ASIN-A", "FNSKU-A", "Alpha"],
        ],
    )
    return path


def test_template_aggregation_preserves_identity_titles_aliases_and_order(
    source: Path,
) -> None:
    rows = aggregate_skus(source, build_asin_sku_map(source))

    assert [row["canonical_sku"] for row in rows] == ["sku-a", "SKU-Z"]
    assert rows[1] == {
        "canonical_sku": "SKU-Z",
        "asin": "ASIN-Z",
        "reference_fnsku": "FNSKU-Z",
        "product_name_latest_observed": "Café latest",
        "product_name_version_count": 2,
        "source_sku_aliases": "Amazon.Found.ASIN-Z",
        "return_record_count_raw": 3,
        "category_id": "water_shoes",
    }


def test_workbook_preserves_fields_tables_validation_and_styles(
    source: Path, tmp_path: Path
) -> None:
    output = tmp_path / "template.xlsx"
    rows = aggregate_skus(source, build_asin_sku_map(source))
    build_workbook(rows, output)
    verify_workbook(output, expected_rows=2)
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["SKU客观特征", "Listing声明"]
        sheet = workbook["SKU客观特征"]
        assert [cell.value for cell in sheet[1]] == [
            "canonical_sku",
            "asin",
            "reference_fnsku",
            "product_name_latest_observed",
            "product_name_version_count",
            "source_sku_aliases",
            "return_record_count_raw",
            "category_id",
            "product_family_id",
            "size_label",
            "size_system",
            "color",
            "package_configuration",
            "fact_source",
            "fact_updated_at",
            "width_design",
            "toe_box_design",
            "upper_material",
            "outsole_material",
            "insole_material",
            "sole_structure",
            "drainage_structure",
            "closure_style",
            "fact_notes",
        ]
        assert sheet["A2"].value == "sku-a"
        assert sheet["D3"].value == "Café latest"
        assert sheet["G3"].value == 3
        assert sheet.tables["SkuObjectiveFacts"].ref == "A1:X3"
        assert sheet.freeze_panes == "A2"
        assert sheet["A1"].comment is not None
        assert sheet["A1"].fill.fgColor.rgb == "00D9EAF7"
        assert sheet["I1"].fill.fgColor.rgb == "00FFF2CC"
        assert sheet["P1"].fill.fgColor.rgb == "00E2F0D9"
        assert sheet["O2"].number_format == "yyyy-mm-dd"
        assert {str(item.sqref) for item in sheet.data_validations.dataValidation} == {
            "K2:K3",
            "N2:N3",
            "P2:Q3",
        }
        rules = list(sheet.conditional_formatting)
        assert len(rules) == 1 and str(rules[0].sqref) == "I2:O3"
        assert sheet.conditional_formatting[rules[0]][0].formula == ['I2=""']
        claims = workbook["Listing声明"]
        assert claims.tables["ListingClaims"].ref == "A1:I2"
        assert [cell.value for cell in claims[1]] == [
            "product_family_id",
            "canonical_sku",
            "claim_type",
            "claim_text",
            "source_location",
            "effective_from",
            "effective_to",
            "captured_at",
            "claim_notes",
        ]
        assert {str(item.sqref) for item in claims.data_validations.dataValidation} == {
            "C2:C10000",
            "E2:E10000",
        }
    finally:
        workbook.close()


@pytest.mark.parametrize(
    "rows, message",
    [
        ([["2026-01-01", "Amazon.Found.A", "A", "A", "Alias"]], "没有规范 SKU"),
        ([["2026-01-01", "SKU", "A", "A", "Title"]], "无法唯一确定参考 FNSKU"),
        (
            [
                ["2026-01-01", "SKU", "A", "F", "Title"],
                ["2026-01-01", "OTHER", "A", "F", "Title"],
            ],
            "无法唯一映射规范 SKU",
        ),
    ],
)
def test_invalid_identity_does_not_produce_template_rows(
    tmp_path: Path, rows: list[list[str]], message: str
) -> None:
    path = tmp_path / "invalid.csv"
    _write_source(path, rows)
    with pytest.raises(ValueError, match=message):
        aggregate_skus(path, build_asin_sku_map(path))


def test_template_command_works_from_another_directory(
    source: Path, tmp_path: Path
) -> None:
    script = (
        Path(__file__).parents[1] / "scripts/generate_water_shoe_feature_template.py"
    )
    output = tmp_path / "result.xlsx"
    result = subprocess.run(
        [
            sys.executable,
            "-X",
            "utf8",
            "-I",
            str(script),
            "--source",
            str(source),
            "--output",
            str(output),
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=True,
    )
    assert "规范 SKU 数: 2" in result.stdout
    verify_workbook(output, expected_rows=2)
