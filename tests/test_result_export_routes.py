"""导出状态与文件接口的认证和失效原因契约。"""

from io import BytesIO

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_classification_result_download import context as context
from test_result_export_jobs import exports as exports

from web_backend.routers.result_export_routes import create_result_export_router


def test_download_routes_auth_pending_failure_and_refresh(context, exports):
    authenticated = True

    def current_user():
        if not authenticated:
            raise HTTPException(status_code=401, detail="请登录")
        return {"id": "user-1"}

    app = FastAPI()
    app.include_router(create_result_export_router(exports, current_user))
    client = TestClient(app)
    response = client.post(
        "/api/result-exports", json={"version_ids": [context.version["version_id"]]}
    )
    assert response.status_code == 200
    job_id = response.json()["id"]
    base = f"/api/result-exports/{job_id}"
    assert client.get(base + "/file").status_code == 409
    exports.run(job_id)
    assert client.get(base).json()["status"] == "ready"
    response = client.get(base + "/file")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book["源明细"].max_row == 4
    book.close()
    authenticated = False
    for path in [base, base + "/file"]:
        assert client.get(path).status_code == 401
    assert (
        client.post(
            "/api/result-exports", json={"version_ids": [context.version["version_id"]]}
        ).status_code
        == 401
    )
    authenticated = True
    assert client.get("/api/result-exports/invalid").status_code == 410
    assert (
        client.post("/api/result-exports", json={"version_ids": []}).status_code == 422
    )


@pytest.mark.parametrize(
    "state", ["expired", "cleaned", "file_missing", "result_missing"]
)
def test_routes_distinguish_cache_unavailable_from_missing_result(
    context, exports, state
):
    job = exports.create([context.version["version_id"]])
    exports.run(job["id"])
    if state == "expired":
        manifest = exports._load(job["id"])
        manifest["expires_at"] = 0
        exports._save(manifest)
    elif state == "cleaned":
        exports._path(job["id"]).unlink()
    elif state == "file_missing":
        exports._path(job["id"], ".xlsx").unlink()
    else:
        with context.database.transaction() as connection:
            connection.execute("DELETE FROM classification_result_records")
            connection.execute("DELETE FROM classification_units")
            connection.execute("UPDATE task_segments SET result_version_id=NULL")
            connection.execute("DELETE FROM classification_result_versions")
    app = FastAPI()
    app.include_router(
        create_result_export_router(exports, lambda: {"id": "synthetic"})
    )
    client = TestClient(app)
    base = f"/api/result-exports/{job['id']}"
    for endpoint in [base, base + "/file"]:
        assert client.get(endpoint).status_code == (
            404 if state == "result_missing" else 410
        )
    if state != "result_missing":
        assert context.service.get(context.version["version_id"])
        regenerated = client.post(
            "/api/result-exports", json={"version_ids": [context.version["version_id"]]}
        )
        assert regenerated.status_code == 200
        assert regenerated.json()["status"] == "queued"
    else:
        assert (
            client.post(
                "/api/result-exports",
                json={"version_ids": [context.version["version_id"]]},
            ).status_code
            == 404
        )
    operation = app.openapi()["paths"]["/api/result-exports/{job_id}"]["get"]
    assert {"404", "410"} <= operation["responses"].keys()
