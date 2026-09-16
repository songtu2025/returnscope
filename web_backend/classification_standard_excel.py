from __future__ import annotations

import hashlib
from copy import deepcopy
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, Field


class ExcelColumnMapping(BaseModel):
    hierarchy_columns: list[str] = Field(min_length=2)
    source_label_column: str | None = None
    sentiment_column: str | None = None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _node_code(namespace: str, kind: str, path: tuple[str, ...]) -> str:
    identity = "\x1f".join((namespace, kind, *path))
    return f"{kind}_{hashlib.sha256(identity.encode()).hexdigest()[:20].upper()}"


def _existing_paths(content: dict[str, Any]) -> tuple[dict, dict]:
    categories = {item["code"]: item for item in content.get("categories", [])}

    def path(item: dict[str, Any]) -> tuple[str, ...]:
        names = [item["name"]]
        parent = item.get("parent_code")
        visited: set[str] = set()
        while parent in categories and parent not in visited:
            visited.add(parent)
            category = categories[parent]
            names.insert(0, category["name"])
            parent = category.get("parent_code")
        return tuple(names)

    category_paths = {path(item): item for item in categories.values()}
    label_items = (
        content.get("labels", []) if content.get("structure_version") == 2 else []
    )
    label_paths = {path(item): item for item in label_items}
    if len(category_paths) != len(categories) or len(label_paths) != len(label_items):
        raise ValueError("现有草稿中同一路径存在多个编码，请先修复重复节点后再导入")
    return category_paths, label_paths


def _read_rows(sheet: Any) -> list[list[str]]:
    rows = [[_text(cell.value) for cell in row] for row in sheet.iter_rows()]
    # 仅还原 Excel 明确标记的合并范围，不推断普通空白行的归属。
    for merged in sheet.merged_cells.ranges:
        value = rows[merged.min_row - 1][merged.min_col - 1]
        for row in range(merged.min_row - 1, merged.max_row):
            for column in range(merged.min_col - 1, merged.max_col):
                rows[row][column] = value
    return rows


def _sentiments(value: str) -> list[str]:
    return {
        "负向": ["NEGATIVE"],
        "负面": ["NEGATIVE"],
        "NEGATIVE": ["NEGATIVE"],
        "正向": ["POSITIVE"],
        "正面": ["POSITIVE"],
        "POSITIVE": ["POSITIVE"],
        "中性": ["NEUTRAL"],
        "NEUTRAL": ["NEUTRAL"],
    }.get(value.upper(), [])


def _parse_rows(
    rows: list[list[str]],
    mapping: ExcelColumnMapping,
    current: dict[str, Any],
    namespace: str,
    sheet_name: str,
) -> dict[str, Any]:
    headers = rows[0]
    indices = [headers.index(name) for name in mapping.hierarchy_columns]
    source_index = (
        headers.index(mapping.source_label_column)
        if mapping.source_label_column
        else None
    )
    sentiment_index = (
        headers.index(mapping.sentiment_column) if mapping.sentiment_column else None
    )
    old_categories, old_labels = _existing_paths(current)
    categories: dict[tuple[str, ...], dict[str, Any]] = {}
    labels: dict[tuple[str, ...], dict[str, Any]] = {}
    issues: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []
    for number, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        path = tuple(row[index] for index in indices)
        while path and not path[-1]:
            path = path[:-1]
        if len(path) < 2 or any(not name for name in path):
            issues.append(
                {
                    "severity": "blocking",
                    "row": number,
                    "message": "至少需要分类和末端标签，且不能缺失中间层级",
                }
            )
            continue
        sentiments = (
            _sentiments(row[sentiment_index]) if sentiment_index is not None else []
        )
        if not sentiments:
            issues.append(
                {
                    "severity": "warning",
                    "row": number,
                    "message": "评价方向为空或尚未明确，已保留路径与来源；请在草稿中选择方向后再发布",
                }
            )
        parent = None
        for depth in range(1, len(path)):
            prefix = path[:depth]
            categories.setdefault(
                prefix,
                {
                    "code": old_categories.get(prefix, {}).get("code")
                    or _node_code(namespace, "CAT", prefix),
                    "name": prefix[-1],
                    "parent_code": parent,
                },
            )
            parent = categories[prefix]["code"]
        existing = labels.get(path)
        if existing and existing["allowed_sentiments"] != sentiments:
            issues.append(
                {
                    "severity": "warning"
                    if not sentiments or not existing["allowed_sentiments"]
                    else "blocking",
                    "row": number,
                    "message": "同一路径的评价方向冲突或未明确，已清空该标签方向，请核对所有来源后选择",
                }
            )
            existing["allowed_sentiments"] = []
        if existing is None:
            existing = deepcopy(old_labels.get(path)) or {
                "code": _node_code(namespace, "LABEL", path),
                "description": "",
                "keywords": [],
                "exclusions": [],
                "examples": [],
                "allowed_claim_ids": [],
            }
            existing.update(
                name=path[-1],
                parent_code=parent,
                group=path[0],
                allowed_sentiments=sentiments,
            )
            labels[path] = existing
        sources.append(
            {
                "sheet": sheet_name,
                "row": number,
                "path": list(path),
                "label_code": existing["code"],
                "source_label": row[source_index] if source_index is not None else "",
                "source_sentiment": row[sentiment_index]
                if sentiment_index is not None
                else "",
            }
        )
    return _build_preview(current, categories, labels, sources, issues)


def _build_preview(
    current: dict[str, Any],
    categories: dict[tuple[str, ...], dict[str, Any]],
    labels: dict[tuple[str, ...], dict[str, Any]],
    sources: list[dict[str, Any]],
    issues: list[dict[str, Any]],
) -> dict[str, Any]:
    content = deepcopy(current)
    content.update(
        structure_version=2,
        categories=list(categories.values()),
        labels=list(labels.values()),
        import_sources=sources,
    )
    content.setdefault("validation_rules", {})["allowed_groups"] = list(
        dict.fromkeys(path[0] for path in labels)
    )
    if labels.keys() & categories.keys():
        issues.append(
            {
                "severity": "blocking",
                "row": None,
                "message": "同一路径同时是分类节点和末端标签，请确认不同深度行的含义",
            }
        )
    names = [path[-1] for path in labels]
    if len(names) != len(set(names)):
        issues.append(
            {
                "severity": "warning",
                "row": None,
                "message": "存在分属不同路径的同名标签，已分别保留，请核对。",
            }
        )
    previous_categories, previous_labels = _existing_paths(current)
    previous_paths = {
        node["code"]: path
        for path, node in {**previous_categories, **previous_labels}.items()
    }
    if any(
        node["code"] in previous_paths and previous_paths[node["code"]] != path
        for path, node in {**categories, **labels}.items()
    ):
        issues.append(
            {
                "severity": "blocking",
                "row": None,
                "message": "导入路径使用了已改名或移动节点的编码，请先确认节点对应关系，不能自动恢复旧路径。",
            }
        )
    return {
        "content": content,
        "issues": issues,
        "stats": {
            "rows": len(sources),
            "categories": len(categories),
            "labels": len(labels),
        },
    }


def preview_excel(
    data: bytes,
    sheet_name: str,
    columns: dict[str, Any],
    current: dict[str, Any],
    namespace: str,
) -> dict[str, Any]:
    """将用户明确选择的工作表解析为草稿预览，不写入业务数据。"""
    try:
        workbook = load_workbook(BytesIO(data), data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ValueError("无法读取 Excel 文件，请上传有效的 .xlsx 文件") from exc
    try:
        result: dict[str, Any] = {
            "sheets": workbook.sheetnames,
            "headers": [],
            "content": None,
            "issues": [],
            "stats": {},
        }
        if not sheet_name:
            return result
        if sheet_name not in workbook.sheetnames:
            raise ValueError("选择的工作表不存在")
        rows = _read_rows(workbook[sheet_name])
        if not rows:
            raise ValueError("工作表为空")
        result["headers"] = rows[0]
        if not columns:
            return result
        mapping = ExcelColumnMapping.model_validate(columns)
        selected = [
            *mapping.hierarchy_columns,
            mapping.source_label_column,
            mapping.sentiment_column,
        ]
        for name in filter(None, selected):
            if rows[0].count(name) != 1:
                raise ValueError(f"列名必须存在且唯一：{name}")
        if len(set(mapping.hierarchy_columns)) != len(mapping.hierarchy_columns):
            raise ValueError("层级列不能重复")
        result.update(_parse_rows(rows, mapping, current, namespace, sheet_name))
        return result
    finally:
        workbook.close()
