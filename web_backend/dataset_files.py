from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from return_semantics.data import (
    PRODUCT_CATEGORY_COLUMNS,
    PRODUCT_COLUMNS,
    PRODUCT_DETAIL_COLUMNS,
    RETURN_COLUMNS,
    RETURN_STORE_COLUMN,
    read_return_file,
)

ALLOWED_EXTENSIONS = {
    "returns": {".csv", ".xlsx"},
    "products": {".xlsx"},
}
PRODUCT_WORKSHEET = "产品信息汇总表"


class DatasetRevisionConflict(ValueError):
    pass


def _identifier_prefix(value: object) -> str:
    match = re.match(r"^([A-Za-z]+\d+)(?:-|$)", str(value or "").strip())
    return match.group(1).casefold() if match else ""


def _product_identity_conflicts(frame: pd.DataFrame) -> pd.DataFrame:
    columns = [*PRODUCT_COLUMNS, *PRODUCT_DETAIL_COLUMNS]
    values = frame.reindex(columns=columns, fill_value="").fillna("")
    listing_prefixes = values["Listing"].map(_identifier_prefix)
    compared_columns = ["MSKU", "产品名称", "SKU"]
    conflict = pd.Series(False, index=values.index)
    for column in compared_columns:
        prefixes = values[column].map(_identifier_prefix)
        conflict |= (
            listing_prefixes.ne("") & prefixes.ne("") & prefixes.ne(listing_prefixes)
        )
    return values.loc[conflict]


def _inspect_returns(path: Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    frame = read_return_file(path)
    missing = [column for column in RETURN_COLUMNS if column not in frame.columns]
    if missing:
        raise ValueError(f"退货数据缺少字段：{', '.join(missing)}")
    valid_comments = int(frame["customer-comments"].fillna("").str.strip().ne("").sum())
    stores = (
        frame[RETURN_STORE_COLUMN].fillna("").astype(str).str.strip()
        if RETURN_STORE_COLUMN in frame.columns
        else pd.Series("", index=frame.index, dtype=str)
    )
    skus = frame["sku"].fillna("").astype(str).str.strip()
    comments = frame["customer-comments"].fillna("").astype(str)
    encoding_anomaly = comments.str.contains(
        r"(?:[A-Za-z][一-鿿]|[一-鿿][A-Za-z])",
        regex=True,
    )
    encoding_anomaly_rows = int(encoding_anomaly.sum())
    matching_key_ready = stores.ne("") & skus.ne("")
    matching_key_ready_rows = int(matching_key_ready.sum())
    quality = {
        "required_columns": len(RETURN_COLUMNS),
        "missing_required_columns": [],
        "valid_comment_rows": valid_comments,
        "valid_comment_rate": round(valid_comments / max(len(frame), 1) * 100, 2),
        "store_column_present": RETURN_STORE_COLUMN in frame.columns,
        "missing_store_rows": int(stores.eq("").sum()),
        "missing_sku_rows": int(skus.eq("").sum()),
        "matching_key_ready_rows": matching_key_ready_rows,
        "matching_key_ready_rate": round(
            matching_key_ready_rows / max(len(frame), 1) * 100,
            2,
        ),
        "stores": sorted(value for value in stores.unique().tolist() if value),
        "text_encoding_anomaly_rows": encoding_anomaly_rows,
        "text_encoding_anomaly_rate": round(
            encoding_anomaly_rows / max(len(frame), 1) * 100,
            2,
        ),
        "text_encoding_anomaly_examples": comments.loc[encoding_anomaly]
        .head(5)
        .tolist(),
    }
    return frame, quality


def _fill_missing_return_store(path: Path, default_store: str) -> None:
    clean_store = default_store.strip()
    if not clean_store:
        return
    frame = read_return_file(path)
    if path.suffix.lower() == ".xlsx":
        _fill_missing_return_store_xlsx(
            path,
            str(frame.attrs["worksheet_name"]),
            clean_store,
        )
        return
    if RETURN_STORE_COLUMN not in frame.columns:
        frame[RETURN_STORE_COLUMN] = clean_store
    else:
        stores = frame[RETURN_STORE_COLUMN].fillna("").astype(str).str.strip()
        frame.loc[stores.eq(""), RETURN_STORE_COLUMN] = clean_store
    frame.to_csv(path, index=False, encoding="utf-8-sig")


def _fill_missing_return_store_xlsx(
    path: Path,
    sheet_name: str,
    default_store: str,
) -> None:
    workbook = load_workbook(path)
    try:
        sheet = workbook[sheet_name]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        if RETURN_STORE_COLUMN in headers:
            column_index = headers.index(RETURN_STORE_COLUMN) + 1
            sheet.cell(row=1, column=column_index).value = RETURN_STORE_COLUMN
        else:
            column_index = sheet.max_column + 1
            sheet.cell(row=1, column=column_index).value = RETURN_STORE_COLUMN
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            if not str(cell.value or "").strip():
                cell.value = default_store
        workbook.save(path)
    finally:
        workbook.close()


def _inspect_products(path: Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    try:
        frame = pd.read_excel(path, sheet_name=PRODUCT_WORKSHEET, dtype=str)
    except ValueError as exc:
        raise ValueError("商品维度缺少“产品信息汇总表”工作表") from exc
    missing = [column for column in PRODUCT_COLUMNS if column not in frame.columns]
    if missing:
        raise ValueError(f"商品维度缺少字段：{', '.join(missing)}")
    complete = int(
        frame[PRODUCT_COLUMNS]
        .fillna("")
        .apply(lambda column: column.str.strip().ne(""))
        .all(axis=1)
        .sum()
    )
    missing_category_columns = [
        column for column in PRODUCT_CATEGORY_COLUMNS if column not in frame.columns
    ]
    if missing_category_columns:
        category_ready = pd.Series(False, index=frame.index)
    else:
        category_ready = (
            frame[PRODUCT_CATEGORY_COLUMNS]
            .fillna("")
            .apply(lambda column: column.str.strip().ne(""))
            .all(axis=1)
        )
    category_ready_rows = int(category_ready.sum())
    identity_conflicts = _product_identity_conflicts(frame)
    identity_conflict_rows = len(identity_conflicts)
    quality = {
        "required_columns": len(PRODUCT_COLUMNS),
        "missing_required_columns": [],
        "complete_rows": complete,
        "complete_rate": round(complete / max(len(frame), 1) * 100, 2),
        "missing_category_columns": missing_category_columns,
        "category_ready_rows": category_ready_rows,
        "category_ready_rate": round(
            category_ready_rows / max(len(frame), 1) * 100,
            2,
        ),
        "missing_category_rows": len(frame) - category_ready_rows,
        "product_identity_conflict_rows": identity_conflict_rows,
        "product_identity_conflict_rate": round(
            identity_conflict_rows / max(len(frame), 1) * 100,
            2,
        ),
        "product_identity_conflict_examples": [
            {
                "listing": str(row["Listing"]),
                "product_name": str(row["产品名称"]),
                "product_sku": str(row["SKU"]),
            }
            for row in identity_conflicts.head(5).to_dict(orient="records")
        ],
    }
    return frame, quality


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _inspect_file_with_frame(
    path: Path,
    kind: str,
) -> tuple[pd.DataFrame, int, int, list[dict[str, str]], dict[str, Any]]:
    if path.suffix.lower() not in ALLOWED_EXTENSIONS.get(kind, set()):
        allowed = "、".join(sorted(ALLOWED_EXTENSIONS.get(kind, set())))
        raise ValueError(f"{kind} 仅支持 {allowed} 文件")
    if kind == "returns":
        frame, quality = _inspect_returns(path)
    elif kind == "products":
        frame, quality = _inspect_products(path)
    else:
        raise ValueError("未知数据类型")
    schema = [
        {"name": str(column), "type": str(frame[column].dtype)}
        for column in frame.columns
    ]
    return frame, len(frame), len(frame.columns), schema, quality


def inspect_file(
    path: Path,
    kind: str,
) -> tuple[int, int, list[dict[str, str]], dict[str, Any]]:
    _, row_count, column_count, schema, quality = _inspect_file_with_frame(path, kind)
    return row_count, column_count, schema, quality


def _return_source_key(stores: list[str]) -> str:
    normalized = sorted(
        {
            re.sub(r"\s+", "", str(value)).upper()
            for value in stores
            if str(value).strip()
        }
    )
    return "|".join(normalized)


def _return_source_name(stores: list[str], original_name: str) -> str:
    if stores:
        labels = [re.sub(r"[:_/\\-]+", " ", value).strip() for value in stores]
        return f"{'、'.join(labels)} 退货数据"
    stem = re.sub(r"[_-]+", " ", Path(original_name).stem).strip()
    return f"{stem or '未命名'} 退货数据"
