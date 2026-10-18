from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from importlib.metadata import version
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles.styleable import StyleDescriptor
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

from web_backend.classification_results import result_workbook as renderer


def test_styles_keep_text_dates_stripes_and_sheet_layout():
    values = [None, 0, False, date(2026, 10, 17), datetime(2026, 10, 17, 12, 30)]
    values += [
        "=SUM(A1)",
        "+公式样文本",
        "中" * 30,
        "中文 English 😀\n第二行",
        "长" * 800,
    ]
    rows = [{"产品名称": value, "反馈正文": value} for value in values]
    content = renderer.result_workbook(rows, [], rows)
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ["反馈汇总", "语义明细", "源明细"]
    for sheet in book:
        assert sheet.freeze_panes == ("D2" if sheet.title == "反馈汇总" else "C2")
        assert sheet.sheet_view.zoomScale == 85
        assert sheet.sheet_view.showGridLines is False
        assert sheet.row_dimensions[1].height == 32
        assert sheet.auto_filter.ref == sheet.dimensions
        assert sheet["A1"].font.bold
        assert sheet["A1"].font.color.rgb == "00FFFFFF"
        assert sheet["A1"].fill.fgColor.rgb == "00176C57"
    for title in ["反馈汇总", "源明细"]:
        sheet = book[title]
        headers = list(next(sheet.values))
        column = headers.index("产品名称") + 1
        for index, value in enumerate(values, 2):
            cell = sheet.cell(index, column)
            expected = (
                datetime.combine(value, datetime.min.time())
                if type(value) is date
                else value
            )
            assert cell.value == expected
            assert cell.font.name == "微软雅黑"
            assert cell.font.sz == 11
            assert cell.font.color.rgb == "00153D34"
            assert not cell.font.bold
            assert cell.alignment.vertical == "top"
            assert cell.alignment.wrap_text
            assert cell.fill.fgColor.rgb == (
                "00F0F7F4" if index % 2 == 0 else "00000000"
            )
            if isinstance(value, str):
                assert cell.data_type == "s"
            if isinstance(value, date):
                assert cell.is_date
                assert cell.number_format == (
                    "yyyy-mm-dd" if type(value) is date else "yyyy-mm-dd h:mm:ss"
                )
        assert sheet.row_dimensions[9].height == 58
        assert sheet.row_dimensions[11].height == 180
    book.close()


def test_repeated_cells_do_not_repeat_style_registration(monkeypatch):
    assignments = []
    original = StyleDescriptor.__set__

    def counted(descriptor, instance, value):
        if isinstance(instance, Cell):
            assignments.append(descriptor.key)
        return original(descriptor, instance, value)

    monkeypatch.setattr(StyleDescriptor, "__set__", counted)
    rows = [{"产品名称": "相同产品", "反馈正文": "相同正文"}] * 100
    book = load_workbook(BytesIO(renderer.result_workbook(rows, rows, rows)))
    assert len(assignments) < 20
    assert all(sheet.max_row == 101 for sheet in book)
    book.close()


def test_height_cache_keeps_width_and_is_limited_to_one_export(monkeypatch):
    calls = Counter()
    original = renderer._display_lines

    def counted(value, width):
        calls[value, width] += 1
        return original(value, width)

    monkeypatch.setattr(renderer, "_display_lines", counted)
    rows = [{"产品名称": "中" * 30, "反馈正文": "中" * 30}] * 20
    renderer.result_workbook(rows, [], rows)
    assert calls["中" * 30, 28] == calls["中" * 30, 54] == 1
    renderer.result_workbook(rows, [], rows)
    assert calls["中" * 30, 28] == calls["中" * 30, 54] == 2


@pytest.mark.parametrize("empty", [False, True])
def test_prepared_cells_keep_original_workbook_xml(monkeypatch, empty):
    assert version("openpyxl") == "3.1.5"
    values = [
        None,
        "",
        0,
        False,
        2.5,
        date(2026, 10, 18),
        datetime(2026, 10, 18, 0, 30),
        "00123",
        "=SUM(A1)",
        "#DIV/0!",
        " 中文 English 😀 <>&\n第二行 ",
        "长" * 800,
    ]
    rows = (
        []
        if empty
        else [
            dict.fromkeys(columns, value)
            for value in values
            for columns in [renderer.SUMMARY_COLUMNS, renderer.SOURCE_COLUMNS]
        ]
    )
    contents = [renderer.result_workbook(rows, rows, rows)]
    monkeypatch.setattr(renderer, "_ResultWorksheet", WriteOnlyWorksheet)
    contents.append(renderer.result_workbook(rows, rows, rows))
    with ZipFile(BytesIO(contents[0])) as new, ZipFile(BytesIO(contents[1])) as old:
        assert new.namelist() == old.namelist()
        for name in new.namelist():
            left, right = new.read(name), old.read(name)
            if name == "docProps/core.xml":
                trees = [ElementTree.fromstring(left), ElementTree.fromstring(right)]
                for tree in trees:
                    for child in list(tree):
                        if child.tag.rsplit("}", 1)[-1] in {"created", "modified"}:
                            tree.remove(child)
                left, right = [ElementTree.tostring(tree) for tree in trees]
            assert left == right, name


def test_prepared_cells_bind_once_without_changing_other_workbooks(monkeypatch):
    calls = []
    original = Cell._bind_value

    def counted(cell, value):
        calls.append(value)
        return original(cell, value)

    monkeypatch.setattr(Cell, "_bind_value", counted)
    rows = [{"产品名称": "合成产品", "反馈正文": "=原文\n第二行"}] * 20
    book = load_workbook(
        BytesIO(renderer.result_workbook(rows, [], rows)), read_only=True
    )
    assert not any(isinstance(value, Cell) for value in calls)
    assert sum(value == "合成产品" for value in calls) == 40
    assert list(book["反馈汇总"].values)[1][2:4] == ("合成产品", "=原文\n第二行")
    book.close()
    normal = Workbook(write_only=True)
    sheet = normal.create_sheet("普通工作表")
    assert type(sheet) is WriteOnlyWorksheet
    sheet.append(["普通文本", 42])
    output = BytesIO()
    normal.save(output)
    normal.close()
    reopened = load_workbook(output, read_only=True)
    assert list(reopened.active.values) == [("普通文本", 42)]
    reopened.close()
