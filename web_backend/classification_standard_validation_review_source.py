from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook

from web_backend.classification_validation_quality import append_reference_fact


class ClassificationStandardValidationReviewSourceMixin:
    if TYPE_CHECKING:

        def _published_config_id(self) -> str: ...

        @staticmethod
        def _round_robin_samples(
            items: list[dict[str, Any]],
            sample_size: int,
            *,
            bucket_fields: tuple[str, ...],
        ) -> list[dict[str, Any]]: ...

    def _review_source_context(
        self,
        filename: str,
        content: bytes,
        draft: dict[str, Any],
        sample_size: int,
    ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        if not filename.lower().endswith(".xlsx"):
            raise ValueError("Review 验证目前支持 .xlsx 文件")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("无法读取 Review 表格，请检查文件格式") from exc
        candidates: list[dict[str, Any]] = []
        skipped = 0
        try:
            references = self._read_references(workbook, draft["snapshot"]["taxonomy"])
            sheet_context = self._review_sheet_context(workbook)
            if sheet_context is not None:
                candidates, skipped = self._review_candidates(
                    *sheet_context,
                    draft["snapshot"]["variants"],
                    references,
                )
        finally:
            workbook.close()
        self._validate_review_references(candidates, references)
        if not candidates:
            raise ValueError(
                "未找到当前品类的评论；表格需包含评论内容列，品类需与当前标准匹配"
            )
        source_id = "review-" + hashlib.sha256(content).hexdigest()
        public = {
            "result_version_id": source_id,
            "source_kind": "review_file",
            "analysis_context": "review",
            "filename": Path(filename).name,
            "available_sample_count": len(candidates),
            "skipped_category_count": skipped,
            "comparison_mode": "draft_only"
            if draft["is_new"]
            else "baseline_and_draft",
            "listing": "Review 样本",
            "sampling": "按店铺及ASIN分层，哈希排序抽样",
        }
        return (
            {
                "kind": "review_file",
                "analysis_context": "review",
                "result": public,
                "config_version_id": self._published_config_id(),
                "standard_key": draft["standard_key"],
                "model_policy_version": draft["snapshot"]["model_policy"]["version"],
            },
            self._round_robin_samples(
                candidates, sample_size, bucket_fields=("store", "listing")
            ),
        )

    @staticmethod
    def _review_sheet_context(workbook: Any) -> tuple[Any, list[str], int] | None:
        for sheet in workbook:
            if sheet.title == "人工参考答案":
                continue
            for number, values in enumerate(
                sheet.iter_rows(max_row=50, values_only=True), start=1
            ):
                if number <= 50 and "评论内容" in values:
                    headers = [str(value or "").strip() for value in values]
                    return sheet, headers, number
        return None

    @staticmethod
    def _review_candidates(
        sheet: Any,
        headers: list[str],
        header_row: int,
        variants: list[dict[str, Any]],
        references: dict[str, Any],
    ) -> tuple[list[dict[str, Any]], int]:
        candidates: list[dict[str, Any]] = []
        seen = set()
        skipped = 0
        sorted_variants = sorted(variants, key=lambda item: -len(item["category_b"]))
        for number, values in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            row = dict(zip(headers, values, strict=False))
            comment = "\n".join(
                str(row.get(key) or "").strip() for key in ("评论标题", "评论内容")
            ).strip()
            if not comment:
                continue
            category = str(row.get("一级品类") or row.get("品类") or "").strip()
            matched = next(
                (
                    variant
                    for variant in sorted_variants
                    if variant["category_b"] in category
                    or variant["category_a"] == category
                ),
                None,
            )
            if category and matched is None:
                skipped += 1
                continue
            identity = str(row.get("评论编号") or number)
            store = str(row.get("下单店铺") or row.get("上架店铺") or "")
            listing = str(row.get("ASIN") or row.get("Listing") or "")
            key = hashlib.sha256(
                f"{identity}\x1f{store}\x1f{listing}\x1f{comment}".encode()
            ).hexdigest()
            if key in seen:
                continue
            seen.add(key)
            candidates.append(
                {
                    "classification_key": key,
                    "comment": comment,
                    "reason": "",
                    "store": store,
                    "listing": listing,
                    "category_a": matched["category_a"] if matched else "",
                    "category_b": matched["category_b"] if matched else "",
                    "source_row": number,
                    "source_category": category,
                    "review_id": identity,
                    "reference": references.get(identity),
                    "record_count": 1,
                    "baseline": {},
                }
            )
        return candidates, skipped

    @staticmethod
    def _validate_review_references(
        candidates: list[dict[str, Any]], references: dict[str, Any]
    ) -> None:
        for identity in references:
            matched_samples = [
                item for item in candidates if item["review_id"] == identity
            ]
            if len(matched_samples) != 1:
                raise ValueError(f"参考答案评论编号 {identity} 未唯一匹配当前品类评论")
            for unit in (
                *references[identity]["units"],
                *references[identity].get("facts", []),
            ):
                if (
                    unit.get("evidence")
                    and unit["evidence"] not in matched_samples[0]["comment"]
                ):
                    raise ValueError(f"参考答案 {identity} 的证据不在原评论中")

    @staticmethod
    def _read_references(workbook, taxonomy: dict) -> dict:
        if "人工参考答案" not in workbook.sheetnames:
            return {}
        rows = iter(workbook["人工参考答案"].values)
        headers = [str(value or "").strip() for value in next(rows, ())]
        required = {"评论编号", "标签编码", "评价方向", "部位", "证据"}
        if not required.issubset(headers):
            raise ValueError(
                "人工参考答案表需包含评论编号、标签编码、评价方向、部位和证据列"
            )
        labels = {item["code"]: item for item in taxonomy["labels"]}
        references: dict[str, dict[str, Any]] = {}
        for values in rows:
            row = dict(zip(headers, values, strict=False))
            identity = str(row.get("评论编号") or "").strip()
            if not identity:
                continue
            reference = references.setdefault(
                identity, {"units": [], "ambiguous": False}
            )
            reference["ambiguous"] |= str(row.get("存在歧义") or "").strip() in {
                "是",
                "1",
                "True",
            }
            code = str(row.get("标签编码") or "").strip()
            append_reference_fact(reference, row, identity, code)
            if code == "无标签":
                continue
            sentiment = str(row.get("评价方向") or "").strip()
            sentiment = {"正向": "POSITIVE", "负向": "NEGATIVE", "中性": "NEUTRAL"}.get(
                sentiment, sentiment
            )
            part = str(row.get("部位") or "UNSPECIFIED").strip()
            if (
                code not in labels
                or sentiment not in labels[code]["allowed_sentiments"]
            ):
                raise ValueError(f"参考答案 {identity} 的标签或评价方向无效")
            if part not in taxonomy["allowed_parts"]:
                raise ValueError(f"参考答案 {identity} 的部位无效")
            evidence = str(row.get("证据") or "").strip()
            if not evidence:
                raise ValueError(f"参考答案 {identity} 缺少证据")
            reference["units"].append(
                {
                    "label_code": code,
                    "sentiment": sentiment,
                    "part": part,
                    "evidence": evidence,
                }
            )
        return references
