from __future__ import annotations

from typing import cast

from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

SKU_HEADERS = [
    "canonical_sku",
    "asin",
    "reference_fnsku",
    "product_name_latest_observed",
    "product_name_version_count",
    "source_sku_aliases",
    "return_record_count_raw",
    "category_id",
    "product_family_id",
    "size_label",
    "size_system",
    "color",
    "package_configuration",
    "fact_source",
    "fact_updated_at",
    "width_design",
    "toe_box_design",
    "upper_material",
    "outsole_material",
    "insole_material",
    "sole_structure",
    "drainage_structure",
    "closure_style",
    "fact_notes",
]


SKU_HEADER_COMMENTS = {
    "canonical_sku": "程序生成的规范 SKU，也是产品特征关联主键，请勿修改。",
    "asin": "当前退货报告中与规范 SKU 唯一对应的 ASIN，请勿修改。",
    "reference_fnsku": "排除 FNSKU=ASIN 等占位值后的参考 FNSKU，请勿修改。",
    "product_name_latest_observed": "退货报告中最近观察到的历史标题，仅用于识别产品，不能作为客观特征。",
    "product_name_version_count": "该 SKU 在当前报告中出现过的标题版本数。",
    "source_sku_aliases": "Amazon.Found.<ASIN> 等原始占位 SKU，供追溯使用。",
    "return_record_count_raw": "当前报告中关联到该 SKU 的原始记录数，包含完全重复记录。",
    "category_id": "首版固定为 water_shoes。",
    "product_family_id": "必填。相同材料、结构和 Listing 声明的一组 SKU 使用同一个产品型号编号。",
    "size_label": "必填。Listing 上展示给买家的完整尺码文字，不要只从 SKU 猜测。",
    "size_system": "必填。选择 US_UNISEX、US_WOMEN、US_MEN、US_KIDS、EU、UK 或 OTHER。",
    "color": "必填。Listing 上展示给买家的颜色名称。",
    "package_configuration": "必填。例如 1 pair / 2 shoes。",
    "fact_source": "必填。客观特征依据，例如 product_spec 或 sample_measurement。",
    "fact_updated_at": "必填。当前客观特征最后确认日期。",
    "width_design": "可选。仅填写可验证的楦型宽度设计，不填写 wide fit 等营销结论。",
    "toe_box_design": "可选。仅填写可验证的鞋头结构宽度。",
    "upper_material": "可选。鞋面客观材质。",
    "outsole_material": "可选。外底客观材质。",
    "insole_material": "可选。鞋垫客观材质。",
    "sole_structure": "可选。鞋底厚度、纹路或分层结构等客观描述。",
    "drainage_structure": "可选。排水孔、导流槽等客观结构；不要在这里填写 quick dry。",
    "closure_style": "可选。例如 slip_on、lace、hook_loop。",
    "fact_notes": "可选。只记录产品事实补充，不记录客户投诉。",
}


CLAIM_HEADERS = [
    "product_family_id",
    "canonical_sku",
    "claim_type",
    "claim_text",
    "source_location",
    "effective_from",
    "effective_to",
    "captured_at",
    "claim_notes",
]


CLAIM_HEADER_COMMENTS = {
    "product_family_id": "必填。必须与 SKU客观特征 工作表中的 product_family_id 一致。",
    "canonical_sku": "可空。为空表示该声明适用于整个产品型号；有值表示只适用于该 SKU。",
    "claim_type": "必填。声明类型使用下拉枚举。",
    "claim_text": "必填。保留 Listing 中的原始英文声明，不要改写。",
    "source_location": "必填。声明出现的位置，例如 title、bullet、A+、size_chart。",
    "effective_from": "已知时填写声明生效日期；未知可空，但不能据此判断历史 Listing 责任。",
    "effective_to": "当前仍有效可空；历史声明填写失效日期。",
    "captured_at": "必填。你采集这条声明的日期。",
    "claim_notes": "可选。补充声明版本或适用范围。",
}


PREFILLED_FILL = PatternFill("solid", fgColor="D9EAF7")


REQUIRED_FILL = PatternFill("solid", fgColor="FFF2CC")


OPTIONAL_FILL = PatternFill("solid", fgColor="E2F0D9")


HEADER_FONT = Font(bold=True, color="1F1F1F")


MISSING_FILL = PatternFill("solid", fgColor="FCE8E6")


def add_list_validation(
    sheet: Worksheet,
    cell_range: str,
    values: list[str],
) -> None:
    formula = '"' + ",".join(values) + '"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "请从下拉列表中选择有效值"
    validation.errorTitle = "无效值"
    validation.prompt = "请选择下拉列表中的值；未知可留空"
    validation.promptTitle = "填写提示"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def style_sheet(
    sheet: Worksheet,
    headers: list[str],
    comments: dict[str, str],
    prefilled_count: int,
    required_count: int,
) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32

    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.value = header
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.comment = Comment(comments[header], "Codex")

        if column_index <= prefilled_count:
            cell.fill = PREFILLED_FILL
        elif column_index <= prefilled_count + required_count:
            cell.fill = REQUIRED_FILL
        else:
            cell.fill = OPTIONAL_FILL

    widths = {
        "canonical_sku": 34,
        "asin": 16,
        "reference_fnsku": 18,
        "product_name_latest_observed": 72,
        "product_name_version_count": 18,
        "source_sku_aliases": 30,
        "return_record_count_raw": 20,
        "category_id": 16,
        "product_family_id": 22,
        "size_label": 22,
        "size_system": 16,
        "color": 20,
        "package_configuration": 22,
        "fact_source": 22,
        "fact_updated_at": 18,
        "width_design": 18,
        "toe_box_design": 18,
        "upper_material": 24,
        "outsole_material": 24,
        "insole_material": 24,
        "sole_structure": 32,
        "drainage_structure": 32,
        "closure_style": 20,
        "fact_notes": 40,
        "claim_type": 20,
        "claim_text": 60,
        "source_location": 20,
        "effective_from": 18,
        "effective_to": 18,
        "captured_at": 18,
        "claim_notes": 40,
    }
    for index, header in enumerate(headers, start=1):
        letter = cast(Cell, sheet.cell(row=1, column=index)).column_letter
        sheet.column_dimensions[letter].width = widths.get(header, 20)
