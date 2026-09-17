from __future__ import annotations

import hashlib
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook

from return_semantics.data import load_return_dataset_auto
from return_semantics.exporter import REVIEW_STATUSES
from return_semantics.prompt import (
    prompt_version,
    recognition_fingerprint,
    validation_contract_matches,
)
from return_semantics.schemas import ProcessingStatus, TaxonomyConfig
from web_backend.classification_standard_service import (
    ClassificationStandardConflict,
    ClassificationStandardService,
)
from web_backend.classification_standard_validation_leakage import (
    find_taxonomy_sample_leaks,
    format_taxonomy_sample_leaks,
)
from web_backend.classification_validation_quality import (
    FACT_QUALITY_POLICY,
    append_reference_fact,
    evaluate_references,
    publication_quality_gate,
)
from web_backend.common import add_audit, json_text, json_value, new_id
from web_backend.database import Database
from web_backend.security import utc_now

if TYPE_CHECKING:
    from web_backend.agent_runner import AgentRunner


class ClassificationStandardValidationNotFound(ValueError):
    pass


class ClassificationStandardValidationConflict(ValueError):
    pass


class ClassificationStandardValidationService:
    def __init__(
        self,
        database: Database,
        standard_service: ClassificationStandardService,
        runner: AgentRunner,
    ) -> None:
        self.database = database
        self.standard_service = standard_service
        self.runner = runner

    def sources(self, draft_id: str) -> list[dict[str, Any]]:
        draft = self.standard_service.get_draft(draft_id)
        raw_sources = [item["public"] for item in self._raw_source_options()]
        if draft["is_new"]:
            return raw_sources
        with self.database.connect() as connection:
            rows = connection.execute(
                """
                SELECT v.id AS result_version_id, v.version_no,
                       v.quality_status, v.unit_count, v.record_count,
                       v.published_at, r.store_site, r.listing,
                       (
                           SELECT COUNT(*) FROM classification_units unit
                           WHERE unit.result_version_id = v.id
                             AND TRIM(COALESCE(unit.comment, '')) != ''
                             AND unit.quality_status != 'excluded'
                       ) AS available_sample_count
                FROM classification_result_versions v
                JOIN classification_results r ON r.id = v.result_id
                WHERE r.standard_version_id = ?
                  AND v.publish_status = 'published'
                  AND EXISTS (
                      SELECT 1 FROM classification_units unit
                      WHERE unit.result_version_id = v.id
                        AND TRIM(COALESCE(unit.comment, '')) != ''
                        AND unit.quality_status != 'excluded'
                  )
                ORDER BY v.published_at DESC, v.id DESC
                """,
                (draft["base_version_id"],),
            ).fetchall()
        return [*raw_sources, *(dict(row) for row in rows)]

    def list_runs(self, draft_id: str) -> list[dict[str, Any]]:
        self.standard_service.get_draft(draft_id)
        with self.database.connect() as connection:
            rows = connection.execute(
                """
                SELECT * FROM classification_standard_validation_runs
                WHERE draft_id = ?
                ORDER BY created_at DESC, id DESC
                LIMIT 20
                """,
                (draft_id,),
            ).fetchall()
        return [self._serialize(dict(row)) for row in rows]

    def get(self, run_id: str, include_items: bool = True) -> dict[str, Any]:
        with self.database.connect() as connection:
            row = connection.execute(
                """
                SELECT * FROM classification_standard_validation_runs
                WHERE id = ?
                """,
                (run_id,),
            ).fetchone()
        if row is None:
            raise ClassificationStandardValidationNotFound("样本验证任务不存在")
        return self._serialize(dict(row), include_items=include_items)

    def create_run(
        self,
        draft_id: str,
        expected_revision: int,
        source_result_version_id: str,
        sample_size: int,
        actor_id: str,
        review_file: tuple[str, bytes] | None = None,
        comparison_type: str = "standard_version",
    ) -> dict[str, Any]:
        if comparison_type not in {"standard_version", "keyword_ab", "semantic_ab"}:
            raise ValueError("不支持的验证目的")
        if sample_size not in {20, 50, 100}:
            raise ValueError("样本规模必须为20、50或100")
        draft = self.standard_service.get_draft(draft_id)
        if int(draft["revision"]) != expected_revision:
            raise ClassificationStandardConflict("草稿已被修改，请刷新后重试")
        if draft["validation"]["blocking"]:
            detail = "；".join(draft["validation"]["blocking"][:3])
            raise ValueError(f"草稿必须先通过结构检查：{detail}")
        source, samples, source_result_version_id = self._validation_source_context(
            draft,
            source_result_version_id,
            sample_size,
            review_file,
        )
        if not samples:
            raise ValueError("所选数据中没有当前品类可用于验证的评论")
        leakage_issues = find_taxonomy_sample_leaks(
            draft["snapshot"]["taxonomy"],
            samples,
        )
        if leakage_issues:
            raise ValueError(format_taxonomy_sample_leaks(leakage_issues))
        self._apply_recognition_context(
            source,
            draft,
            draft_id,
            expected_revision,
            comparison_type,
        )
        run_id = new_id("classification_standard_validation")
        now = utc_now()
        with self.database.transaction(immediate=True) as connection:
            active = connection.execute(
                """
                SELECT id FROM classification_standard_validation_runs
                WHERE draft_id = ? AND status IN ('queued', 'running')
                LIMIT 1
                """,
                (draft_id,),
            ).fetchone()
            if active is not None:
                raise ClassificationStandardValidationConflict(
                    "当前草稿已有样本验证正在运行"
                )
            connection.execute(
                """
                INSERT INTO classification_standard_validation_runs(
                    id, standard_id, draft_id, draft_revision,
                    base_version_id, source_result_version_id,
                    config_version_id, status, stage, sample_size,
                    snapshot_json, source_json, sample_json,
                    created_by, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, 'queued', 'queued', ?, ?, ?, ?, ?, ?)
                """,
                (
                    run_id,
                    draft["standard_id"],
                    draft_id,
                    expected_revision,
                    draft["base_version_id"],
                    source_result_version_id,
                    source.get("config_version_id")
                    or source["task"]["config_version_id"],
                    len(samples),
                    json_text(draft["snapshot"]),
                    json_text(source),
                    json_text(samples),
                    actor_id,
                    now,
                ),
            )
        add_audit(
            self.database,
            "classification_standard_validation",
            run_id,
            "create",
            actor_id,
            after={
                "draft_id": draft_id,
                "draft_revision": expected_revision,
                "source_result_version_id": source_result_version_id,
                "sample_size": len(samples),
            },
        )
        return self.get(run_id)

    def _validation_source_context(
        self,
        draft: dict[str, Any],
        source_result_version_id: str,
        sample_size: int,
        review_file: tuple[str, bytes] | None,
    ) -> tuple[dict[str, Any], list[dict[str, Any]], str]:
        raw_source_ids = {item["id"] for item in self._raw_source_options()}
        if review_file is not None:
            source, samples = self._review_source_context(
                review_file[0], review_file[1], draft, sample_size
            )
            return source, samples, source["result"]["result_version_id"]
        if source_result_version_id in raw_source_ids:
            source, samples = self._raw_source_context(
                source_result_version_id,
                draft,
                sample_size,
            )
            return source, samples, source_result_version_id
        source = self._source_context(
            source_result_version_id,
            str(draft["base_version_id"]),
        )
        return (
            source,
            self._sample(source_result_version_id, sample_size),
            source_result_version_id,
        )

    def _apply_recognition_context(
        self,
        source: dict[str, Any],
        draft: dict[str, Any],
        draft_id: str,
        expected_revision: int,
        comparison_type: str,
    ) -> None:
        candidate = TaxonomyConfig.model_validate(draft["snapshot"]["taxonomy"])
        baseline = self.standard_service.taxonomy_for_version(
            str(draft["base_version_id"])
        )
        if comparison_type != "standard_version":
            candidate = candidate.model_copy(
                update={
                    "version": f"draft-{draft_id}-r{expected_revision}",
                }
            )
            baseline = candidate.model_copy(
                update={
                    "recognition_profile": "legacy_v3"
                    if comparison_type == "keyword_ab"
                    else "keyword_free_v1",
                }
            )
            candidate = candidate.model_copy(
                update={
                    "recognition_profile": "keyword_free_v1"
                    if comparison_type == "keyword_ab"
                    else "semantic_v1",
                }
            )
        source["comparison_type"] = comparison_type
        source["recognition_contract"] = {
            side: {
                "profile": config.recognition_profile,
                "prompt_version": prompt_version(config),
                "fingerprint": recognition_fingerprint(config),
            }
            for side, config in (("baseline", baseline), ("candidate", candidate))
        }
        source["recognition_taxonomies"] = {
            "baseline": baseline.model_dump(mode="json"),
            "candidate": candidate.model_dump(mode="json"),
        }
        source["result"].update(
            {
                "comparison_type": comparison_type,
                "recognition_contract": source["recognition_contract"],
            }
        )
        if comparison_type != "standard_version":
            source["result"]["comparison_mode"] = "baseline_and_draft"
        if candidate.recognition_profile == "fact_v2":
            source["quality_policy"] = deepcopy(FACT_QUALITY_POLICY)
            source["result"]["quality_policy"] = source["quality_policy"]

    def approve(
        self,
        run_id: str,
        expected_revision: int,
        note: str,
        actor_id: str,
    ) -> dict[str, Any]:
        validation = self.get(run_id)
        if (
            validation["source"].get("comparison_type", "standard_version")
            != "standard_version"
        ):
            raise ValueError("策略对照仅用于诊断，不能代替发布验证")
        if not validation["is_current"] or int(validation["draft_revision"]) != int(
            expected_revision
        ):
            raise ClassificationStandardValidationConflict(
                "验证结果不属于当前草稿修订，请重新运行"
            )
        if validation["status"] != "completed":
            raise ValueError("样本验证尚未完成")
        if int(validation["error_count"]) > 0:
            raise ValueError("样本验证存在模型错误，不能确认通过")
        if not validation["quality_gate"]["passed"]:
            raise ValueError(
                "质量门槛未通过：" + "；".join(validation["quality_gate"]["blocking"])
            )
        approval_note = note.strip()
        if not approval_note:
            raise ValueError("请填写验证结论")
        approved_at = utc_now()
        with self.database.transaction(immediate=True) as connection:
            connection.execute(
                """
                UPDATE classification_standard_validation_runs
                SET approved_by = ?, approved_at = ?, approval_note = ?
                WHERE id = ?
                """,
                (actor_id, approved_at, approval_note, run_id),
            )
        add_audit(
            self.database,
            "classification_standard_validation",
            run_id,
            "approve",
            actor_id,
            after={
                "draft_revision": expected_revision,
                "note": approval_note,
            },
        )
        return self.get(run_id)

    def recover(self) -> None:
        with self.database.transaction(immediate=True) as connection:
            connection.execute(
                """
                UPDATE classification_standard_validation_runs
                SET status = 'queued', stage = 'queued', error = NULL,
                    processed_count = 0, started_at = NULL
                WHERE status = 'running'
                """
            )

    def claim_next(self) -> str | None:
        with self.database.transaction(immediate=True) as connection:
            row = connection.execute(
                """
                SELECT id FROM classification_standard_validation_runs
                WHERE status = 'queued'
                ORDER BY created_at, id
                LIMIT 1
                """
            ).fetchone()
            if row is None:
                return None
            run_id = str(row["id"])
            updated = connection.execute(
                """
                UPDATE classification_standard_validation_runs
                SET status = 'running', stage = 'calling_model',
                    started_at = ?, error = NULL
                WHERE id = ? AND status = 'queued'
                """,
                (utc_now(), run_id),
            )
        return run_id if updated.rowcount == 1 else None

    def run(self, run_id: str) -> None:
        with self.database.connect() as connection:
            row = connection.execute(
                """
                SELECT * FROM classification_standard_validation_runs
                WHERE id = ?
                """,
                (run_id,),
            ).fetchone()
        if row is None or row["status"] != "running":
            return
        validation = dict(row)
        samples = json_value(validation["sample_json"], [])
        source = json_value(validation["source_json"], {})
        snapshot = json_value(validation["snapshot_json"], {})
        taxonomy_data = deepcopy(
            source.get("recognition_taxonomies", {}).get(
                "candidate", snapshot["taxonomy"]
            )
        )
        taxonomy_data["version"] = (
            f"draft-{validation['draft_id']}-r{validation['draft_revision']}"
        )
        try:
            taxonomy = TaxonomyConfig.model_validate(taxonomy_data)

            stage = "calling_model"

            def progress(current: int, total: int) -> None:
                if current == 1 or current == total or current % 5 == 0:
                    with self.database.transaction() as connection:
                        connection.execute(
                            """
                            UPDATE classification_standard_validation_runs
                            SET processed_count = ?, stage = ?
                            WHERE id = ? AND status = 'running'
                            """,
                            (current, stage, run_id),
                        )

            if (
                source.get("kind") in {"raw_dataset", "review_file"}
                or source.get("comparison_type", "standard_version")
                != "standard_version"
            ) and source.get("result", {}).get(
                "comparison_mode"
            ) == "baseline_and_draft":
                base_taxonomy = (
                    TaxonomyConfig.model_validate(
                        source["recognition_taxonomies"]["baseline"]
                    )
                    if source.get("recognition_taxonomies")
                    else self.standard_service.taxonomy_for_version(
                        str(validation["base_version_id"])
                    )
                )
                stage = "comparing_baseline"
                progress(0, len(samples))
                baseline_pipeline = self.runner.classify_taxonomy_sample(
                    taxonomy=base_taxonomy,
                    samples=samples,
                    source=source,
                    progress=progress,
                )
                for sample in samples:
                    key = str(sample["classification_key"])
                    sample["baseline"] = baseline_pipeline.classifications[
                        key
                    ].model_dump(mode="json")
            else:
                baseline_pipeline = None

            stage = "calling_model"
            progress(0, len(samples))
            pipeline = self.runner.classify_taxonomy_sample(
                taxonomy=taxonomy,
                samples=samples,
                source=source,
                progress=progress,
            )
            items = self._comparison_items(samples, pipeline.classifications)
            summary = self._summary(items)
            summary["reference_evaluation"] = self._evaluate_references(items)
            if source.get("comparison_type", "standard_version") == "standard_version":
                summary["reference_evaluation"]["sides"].pop("baseline", None)
            usage = {
                "baseline": (
                    baseline_pipeline.usage if baseline_pipeline is not None else {}
                ),
                "draft": pipeline.usage,
            }
            metrics = {
                "baseline": (
                    baseline_pipeline.request_metrics
                    if baseline_pipeline is not None
                    else {}
                ),
                "draft": pipeline.request_metrics,
            }
            summary["efficiency"] = self._efficiency_summary(
                items,
                usage,
                metrics,
                {
                    "baseline": (
                        baseline_pipeline.model_calls
                        if baseline_pipeline is not None
                        else None
                    ),
                    "draft": pipeline.model_calls,
                },
            )
            model_names = sorted(
                {
                    str(item["draft"]["model_name"])
                    for item in items
                    if item["draft"].get("model_name")
                }
                | {
                    str(result.model_name)
                    for result in (
                        baseline_pipeline.classifications.values()
                        if baseline_pipeline is not None
                        else []
                    )
                    if result.model_name
                }
            )
            with self.database.transaction(immediate=True) as connection:
                connection.execute(
                    """
                    UPDATE classification_standard_validation_runs
                    SET status = 'completed', stage = 'completed',
                        processed_count = sample_size, changed_count = ?,
                        unknown_count = ?, review_count = ?, error_count = ?,
                        result_json = ?, summary_json = ?, usage_json = ?,
                        metrics_json = ?, model_names_json = ?,
                        error = NULL, completed_at = ?
                    WHERE id = ? AND status = 'running'
                    """,
                    (
                        summary["changed_count"],
                        summary["unknown_count"],
                        summary["review_count"],
                        summary["error_count"],
                        json_text(items),
                        json_text(summary),
                        json_text(usage),
                        json_text(metrics),
                        json_text(model_names),
                        utc_now(),
                        run_id,
                    ),
                )
        except Exception as exc:
            with self.database.transaction(immediate=True) as connection:
                connection.execute(
                    """
                    UPDATE classification_standard_validation_runs
                    SET status = 'failed', stage = 'failed', error = ?,
                        completed_at = ?
                    WHERE id = ? AND status = 'running'
                    """,
                    (str(exc)[:2000], utc_now(), run_id),
                )

    def _raw_source_options(self) -> list[dict[str, Any]]:
        with self.database.connect() as connection:
            rows = connection.execute(
                """
                SELECT v.id AS version_id, v.version, v.file_path,
                       v.row_count, v.created_at, d.name AS dataset_name,
                       d.kind
                FROM datasets d
                JOIN dataset_versions v
                  ON v.dataset_id = d.id AND v.version = d.current_version
                WHERE d.archived_at IS NULL
                  AND d.kind IN ('returns', 'products')
                ORDER BY v.created_at DESC, v.id
                """
            ).fetchall()
        returns = [dict(row) for row in rows if row["kind"] == "returns"]
        products = [dict(row) for row in rows if row["kind"] == "products"]
        output = []
        for return_version in returns:
            for product_version in products:
                source_key = "\x1f".join(
                    [return_version["version_id"], product_version["version_id"]]
                )
                source_id = (
                    "raw_validation_"
                    + hashlib.sha256(source_key.encode("utf-8")).hexdigest()[:24]
                )
                output.append(
                    {
                        "id": source_id,
                        "return": return_version,
                        "product": product_version,
                        "public": {
                            "result_version_id": source_id,
                            "source_kind": "raw_dataset",
                            "version_no": int(return_version["version"]),
                            "quality_status": "raw",
                            "unit_count": 0,
                            "record_count": int(return_version["row_count"]),
                            "published_at": return_version["created_at"],
                            "store_site": "",
                            "listing": "",
                            "available_sample_count": None,
                            "return_dataset_name": return_version["dataset_name"],
                            "return_version_id": return_version["version_id"],
                            "product_dataset_name": product_version["dataset_name"],
                            "product_version_id": product_version["version_id"],
                        },
                    }
                )
        return output

    def _raw_source_context(
        self,
        source_id: str,
        draft: dict[str, Any],
        sample_size: int,
    ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        option = next(
            (item for item in self._raw_source_options() if item["id"] == source_id),
            None,
        )
        if option is None:
            raise ValueError("所选退货数据或产品信息版本已不可用")
        dataset = load_return_dataset_auto(
            Path(option["return"]["file_path"]),
            Path(option["product"]["file_path"]),
        )
        categories = {
            (str(item["category_a"]), str(item["category_b"]))
            for item in draft["snapshot"]["variants"]
        }
        candidates = []
        for row in dataset.unique_comments.itertuples(index=False):
            category = (str(row.category_a), str(row.category_b))
            if category not in categories or str(row.product_match_status) != "matched":
                continue
            candidates.append(
                {
                    "classification_key": str(row.classification_key),
                    "comment": str(row.comment_normalized),
                    "reason": str(row.reason),
                    "category_a": category[0],
                    "category_b": category[1],
                    "store": str(row.store),
                    "listing": str(row.listing),
                    "record_count": int(row.record_count),
                    "baseline": {},
                }
            )
        samples = self._round_robin_samples(
            candidates,
            sample_size,
            bucket_fields=("store", "listing"),
        )
        if not samples:
            raise ValueError("所选数据中没有当前品类可用于验证的评论")
        config_version_id = self._published_config_id()
        stores = sorted({item["store"] for item in candidates if item["store"]})
        listings = sorted({item["listing"] for item in candidates if item["listing"]})
        public = {
            **option["public"],
            "comparison_mode": (
                "draft_only" if draft["is_new"] else "baseline_and_draft"
            ),
            "unit_count": len(candidates),
            "available_sample_count": len(candidates),
            "store_site": stores[0] if len(stores) == 1 else f"{len(stores)} 个店铺",
            "listing": (
                listings[0] if len(listings) == 1 else f"{len(listings)} 个 Listing"
            ),
        }
        return (
            {
                "kind": "raw_dataset",
                "result": public,
                "config_version_id": config_version_id,
                "standard_key": str(draft["standard_key"]),
                "model_policy_version": str(
                    draft["snapshot"]["model_policy"]["version"]
                ),
                "store": stores[0] if len(stores) == 1 else "",
                "listing": listings[0] if len(listings) == 1 else None,
            },
            samples,
        )

    def _published_config_id(self) -> str:
        with self.database.connect() as connection:
            config = connection.execute(
                """
                SELECT v.id
                FROM api_connections c
                JOIN api_config_versions v ON v.id = c.active_version_id
                WHERE v.published_at IS NOT NULL
                ORDER BY c.updated_at DESC, v.id
                LIMIT 1
                """
            ).fetchone()
        if config is None:
            raise ValueError("请先验证并发布一个模型服务配置")
        return str(config["id"])

    def _review_source_context(
        self,
        filename: str,
        content: bytes,
        draft: dict[str, Any],
        sample_size: int,
    ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        if not filename.lower().endswith(".xlsx"):
            raise ValueError("Review 验证目前支持 .xlsx 文件")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError("无法读取 Review 表格，请检查文件格式") from exc
        candidates = []
        skipped = 0
        try:
            references = self._read_references(workbook, draft["snapshot"]["taxonomy"])
            sheet_context = self._review_sheet_context(workbook)
            if sheet_context is not None:
                candidates, skipped = self._review_candidates(
                    *sheet_context,
                    draft["snapshot"]["variants"],
                    references,
                )
        finally:
            workbook.close()
        self._validate_review_references(candidates, references)
        if not candidates:
            raise ValueError(
                "未找到当前品类的评论；表格需包含评论内容列，品类需与当前标准匹配"
            )
        source_id = "review-" + hashlib.sha256(content).hexdigest()
        public = {
            "result_version_id": source_id,
            "source_kind": "review_file",
            "analysis_context": "review",
            "filename": Path(filename).name,
            "available_sample_count": len(candidates),
            "skipped_category_count": skipped,
            "comparison_mode": "draft_only"
            if draft["is_new"]
            else "baseline_and_draft",
            "listing": "Review 样本",
            "sampling": "按店铺及ASIN分层，哈希排序抽样",
        }
        return (
            {
                "kind": "review_file",
                "analysis_context": "review",
                "result": public,
                "config_version_id": self._published_config_id(),
                "standard_key": draft["standard_key"],
                "model_policy_version": draft["snapshot"]["model_policy"]["version"],
            },
            self._round_robin_samples(
                candidates, sample_size, bucket_fields=("store", "listing")
            ),
        )

    @staticmethod
    def _review_sheet_context(workbook: Any) -> tuple[Any, list[str], int] | None:
        for sheet in workbook:
            if sheet.title == "人工参考答案":
                continue
            for number, values in enumerate(
                sheet.iter_rows(max_row=50, values_only=True), start=1
            ):
                if number <= 50 and "评论内容" in values:
                    headers = [str(value or "").strip() for value in values]
                    return sheet, headers, number
        return None

    @staticmethod
    def _review_candidates(
        sheet: Any,
        headers: list[str],
        header_row: int,
        variants: list[dict[str, Any]],
        references: dict[str, Any],
    ) -> tuple[list[dict[str, Any]], int]:
        candidates = []
        seen = set()
        skipped = 0
        sorted_variants = sorted(variants, key=lambda item: -len(item["category_b"]))
        for number, values in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            row = dict(zip(headers, values, strict=False))
            comment = "\n".join(
                str(row.get(key) or "").strip() for key in ("评论标题", "评论内容")
            ).strip()
            if not comment:
                continue
            category = str(row.get("一级品类") or row.get("品类") or "").strip()
            matched = next(
                (
                    variant
                    for variant in sorted_variants
                    if variant["category_b"] in category
                    or variant["category_a"] == category
                ),
                None,
            )
            if category and matched is None:
                skipped += 1
                continue
            identity = str(row.get("评论编号") or number)
            store = str(row.get("下单店铺") or row.get("上架店铺") or "")
            listing = str(row.get("ASIN") or row.get("Listing") or "")
            key = hashlib.sha256(
                f"{identity}\x1f{store}\x1f{listing}\x1f{comment}".encode()
            ).hexdigest()
            if key in seen:
                continue
            seen.add(key)
            candidates.append(
                {
                    "classification_key": key,
                    "comment": comment,
                    "reason": "",
                    "store": store,
                    "listing": listing,
                    "category_a": matched["category_a"] if matched else "",
                    "category_b": matched["category_b"] if matched else "",
                    "source_row": number,
                    "source_category": category,
                    "review_id": identity,
                    "reference": references.get(identity),
                    "record_count": 1,
                    "baseline": {},
                }
            )
        return candidates, skipped

    @staticmethod
    def _validate_review_references(
        candidates: list[dict[str, Any]], references: dict[str, Any]
    ) -> None:
        for identity in references:
            matched_samples = [
                item for item in candidates if item["review_id"] == identity
            ]
            if len(matched_samples) != 1:
                raise ValueError(f"参考答案评论编号 {identity} 未唯一匹配当前品类评论")
            for unit in (
                *references[identity]["units"],
                *references[identity].get("facts", []),
            ):
                if (
                    unit.get("evidence")
                    and unit["evidence"] not in matched_samples[0]["comment"]
                ):
                    raise ValueError(f"参考答案 {identity} 的证据不在原评论中")

    @staticmethod
    def _read_references(workbook, taxonomy: dict) -> dict:
        if "人工参考答案" not in workbook.sheetnames:
            return {}
        rows = iter(workbook["人工参考答案"].values)
        headers = [str(value or "").strip() for value in next(rows, ())]
        required = {"评论编号", "标签编码", "评价方向", "部位", "证据"}
        if not required.issubset(headers):
            raise ValueError(
                "人工参考答案表需包含评论编号、标签编码、评价方向、部位和证据列"
            )
        labels = {item["code"]: item for item in taxonomy["labels"]}
        references = {}
        for values in rows:
            row = dict(zip(headers, values, strict=False))
            identity = str(row.get("评论编号") or "").strip()
            if not identity:
                continue
            reference = references.setdefault(
                identity, {"units": [], "ambiguous": False}
            )
            reference["ambiguous"] |= str(row.get("存在歧义") or "").strip() in {
                "是",
                "1",
                "True",
            }
            code = str(row.get("标签编码") or "").strip()
            append_reference_fact(reference, row, identity, code)
            if code == "无标签":
                continue
            sentiment = str(row.get("评价方向") or "").strip()
            sentiment = {"正向": "POSITIVE", "负向": "NEGATIVE", "中性": "NEUTRAL"}.get(
                sentiment, sentiment
            )
            part = str(row.get("部位") or "UNSPECIFIED").strip()
            if (
                code not in labels
                or sentiment not in labels[code]["allowed_sentiments"]
            ):
                raise ValueError(f"参考答案 {identity} 的标签或评价方向无效")
            if part not in taxonomy["allowed_parts"]:
                raise ValueError(f"参考答案 {identity} 的部位无效")
            evidence = str(row.get("证据") or "").strip()
            if not evidence:
                raise ValueError(f"参考答案 {identity} 缺少证据")
            reference["units"].append(
                {
                    "label_code": code,
                    "sentiment": sentiment,
                    "part": part,
                    "evidence": evidence,
                }
            )
        return references

    _evaluate_references = staticmethod(evaluate_references)

    @staticmethod
    def _round_robin_samples(
        items: list[dict[str, Any]],
        sample_size: int,
        *,
        bucket_fields: tuple[str, ...],
    ) -> list[dict[str, Any]]:
        buckets: dict[tuple[str, ...], list[dict[str, Any]]] = {}
        for item in sorted(items, key=lambda value: value["classification_key"]):
            key = tuple(str(item[field]) for field in bucket_fields)
            buckets.setdefault(key, []).append(item)
        output = []
        keys = sorted(buckets, key=lambda key: (-len(buckets[key]), key))
        while len(output) < sample_size and any(buckets.values()):
            for key in keys:
                if buckets[key] and len(output) < sample_size:
                    output.append(buckets[key].pop(0))
        return output

    def _source_context(
        self,
        result_version_id: str,
        base_version_id: str,
    ) -> dict[str, Any]:
        with self.database.connect() as connection:
            result = connection.execute(
                """
                SELECT v.id AS result_version_id, v.version_no,
                       v.quality_status, v.unit_count, v.record_count,
                       v.published_at, r.store_site, r.listing,
                       r.standard_version_id,
                       r.source_task_id, r.source_segment_id
                FROM classification_result_versions v
                JOIN classification_results r ON r.id = v.result_id
                WHERE v.id = ? AND v.publish_status = 'published'
                """,
                (result_version_id,),
            ).fetchone()
            if result is None:
                raise ValueError("所选分类结果不存在或尚未发布")
            if result["standard_version_id"] != base_version_id:
                raise ValueError("所选分类结果不属于草稿的基础标准版本")
            task = connection.execute(
                """
                SELECT id, config_version_id, store, listing, snapshot_json
                FROM tasks WHERE id = ?
                """,
                (result["source_task_id"],),
            ).fetchone()
            segment = connection.execute(
                "SELECT * FROM task_segments WHERE id = ?",
                (result["source_segment_id"],),
            ).fetchone()
        if task is None or segment is None:
            raise ValueError("分类结果缺少可复用的任务执行上下文")
        public_result = dict(result)
        public_result.pop("source_task_id", None)
        public_result.pop("source_segment_id", None)
        return {
            "result": public_result,
            "task": dict(task),
            "segment": dict(segment),
        }

    def _sample(
        self,
        result_version_id: str,
        sample_size: int,
    ) -> list[dict[str, Any]]:
        with self.database.connect() as connection:
            rows = connection.execute(
                """
                SELECT unit.classification_key, unit.comment, unit.reason,
                       unit.classification_json, unit.quality_status,
                       MIN(COALESCE(record.category_a, '')) AS category_a,
                       MIN(COALESCE(record.category_b, '')) AS category_b
                FROM classification_units unit
                LEFT JOIN classification_result_records record
                  ON record.result_version_id = unit.result_version_id
                 AND record.classification_key = unit.classification_key
                WHERE unit.result_version_id = ?
                  AND TRIM(COALESCE(unit.comment, '')) != ''
                  AND unit.quality_status != 'excluded'
                GROUP BY unit.classification_key, unit.comment, unit.reason,
                         unit.classification_json, unit.quality_status
                ORDER BY unit.classification_key
                """,
                (result_version_id,),
            ).fetchall()
        buckets: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            item = dict(row)
            classification = json_value(item.pop("classification_json"), {})
            item["baseline"] = classification
            labels = classification.get("primary_label_codes", [])
            bucket = str(labels[0]) if labels else "__NO_PRIMARY_LABEL__"
            buckets.setdefault(bucket, []).append(item)
        output: list[dict[str, Any]] = []
        keys = sorted(buckets, key=lambda key: (-len(buckets[key]), key))
        while len(output) < sample_size and any(buckets.values()):
            for key in keys:
                if buckets[key] and len(output) < sample_size:
                    output.append(buckets[key].pop(0))
        return output

    @staticmethod
    def _comparison_items(
        samples: list[dict[str, Any]],
        classifications: dict[str, Any],
    ) -> list[dict[str, Any]]:
        output = []
        for sample in samples:
            key = str(sample["classification_key"])
            result = classifications[key].model_dump(mode="json")
            baseline_labels = sorted(sample["baseline"].get("primary_label_codes", []))
            draft_labels = sorted(result.get("primary_label_codes", []))

            def signature(units):
                return sorted(
                    {
                        (unit["label_code"], unit.get("sentiment"), unit.get("part"))
                        for unit in units
                    }
                )

            output.append(
                {
                    "classification_key": key,
                    "comment": sample["comment"],
                    "reference": sample.get("reference"),
                    "category_a": sample["category_a"],
                    "category_b": sample["category_b"],
                    "baseline": {
                        "extracted_facts": sample["baseline"].get(
                            "extracted_facts", []
                        ),
                        "fact_mappings": sample["baseline"].get("fact_mappings", []),
                        "primary_label_codes": baseline_labels,
                        "semantic_units": sample["baseline"].get("semantic_units", []),
                        "unknown_semantics": sample["baseline"].get(
                            "unknown_semantics", []
                        ),
                        "status": sample["baseline"].get("status"),
                        "review_reasons": sample["baseline"].get("review_reasons", []),
                        "reason": sample.get("reason"),
                    },
                    "draft": {
                        "extracted_facts": result.get("extracted_facts", []),
                        "fact_mappings": result.get("fact_mappings", []),
                        "primary_label_codes": draft_labels,
                        "status": result["status"],
                        "unknown_semantics": result.get("unknown_semantics", []),
                        "semantic_units": [
                            {
                                "label_code": unit["label_code"],
                                "opinion": unit["opinion"],
                                "evidence": unit["evidence"],
                                "sentiment": unit["sentiment"],
                                "part": unit["part"],
                            }
                            for unit in result.get("semantic_units", [])
                        ],
                        "review_reasons": result.get("review_reasons", []),
                        "model_name": result.get("model_name"),
                    },
                    "changed": (
                        baseline_labels != draft_labels
                        or signature(sample["baseline"].get("semantic_units", []))
                        != signature(result.get("semantic_units", []))
                    ),
                }
            )
        return output

    @staticmethod
    def _summary(items: list[dict[str, Any]]) -> dict[str, Any]:
        total = len(items)
        changed_count = sum(bool(item["changed"]) for item in items)
        unknown_count = sum(bool(item["draft"]["unknown_semantics"]) for item in items)
        review_count = sum(
            str(item["draft"]["status"]) in REVIEW_STATUSES for item in items
        )
        error_count = sum(
            str(item["draft"]["status"]) == ProcessingStatus.MODEL_ERROR.value
            or str(item["baseline"].get("status")) == ProcessingStatus.MODEL_ERROR.value
            for item in items
        )
        coverage_count = sum(bool(item["draft"]["semantic_units"]) for item in items)

        def rate(value: int) -> float:
            return round(value / total * 100, 1) if total else 0.0

        return {
            "sample_size": total,
            "changed_count": changed_count,
            "changed_rate": rate(changed_count),
            "coverage_count": coverage_count,
            "coverage_rate": rate(coverage_count),
            "unknown_count": unknown_count,
            "unknown_rate": rate(unknown_count),
            "review_count": review_count,
            "review_rate": rate(review_count),
            "error_count": error_count,
            "error_rate": rate(error_count),
        }

    @classmethod
    def _efficiency_summary(
        cls,
        items: list[dict[str, Any]],
        usage: dict[str, Any],
        metrics: dict[str, Any],
        model_calls: dict[str, int | None] | None = None,
    ) -> dict[str, Any]:
        sample_size = len(items)
        sides: dict[str, dict[str, int | float]] = {}
        known_model_calls = model_calls or {}
        for side in ("baseline", "draft"):
            side_usage = usage.get(side)
            side_usage = side_usage if isinstance(side_usage, dict) else {}
            side_metrics = metrics.get(side)
            side_metrics = side_metrics if isinstance(side_metrics, dict) else {}
            values = cls._efficiency_side(
                items,
                side,
                sample_size,
                side_usage,
                side_metrics,
                known_model_calls.get(side),
            )
            if values:
                sides[side] = values
        return {"sample_size": sample_size, "sides": sides}

    @staticmethod
    def _efficiency_side(
        items: list[dict[str, Any]],
        side: str,
        sample_size: int,
        usage: dict[str, Any],
        metrics: dict[str, Any],
        model_calls: int | None,
    ) -> dict[str, int | float]:
        values: dict[str, int | float] = {}
        call_count = (
            model_calls if model_calls is not None else metrics.get("fact_model_calls")
        )
        if isinstance(call_count, int) and not isinstance(call_count, bool):
            values["model_calls"] = call_count
            values["average_model_calls"] = (
                round(call_count / sample_size, 2) if sample_size else 0.0
            )
        total_tokens = ClassificationStandardValidationService._total_tokens(usage)
        if total_tokens is not None:
            values["total_tokens"] = total_tokens
            values["average_tokens"] = (
                round(total_tokens / sample_size, 1) if sample_size else 0.0
            )
        coverage_audit_count = metrics.get("coverage_audit_calls")
        if isinstance(coverage_audit_count, int) and not isinstance(
            coverage_audit_count, bool
        ):
            values["coverage_audit_count"] = coverage_audit_count
            values["coverage_audit_rate"] = (
                round(coverage_audit_count / sample_size * 100, 1)
                if sample_size
                else 0.0
            )
        statuses = [
            item.get(side, {}).get("status")
            for item in items
            if item.get(side, {}).get("status")
        ]
        if statuses:
            review_count = sum(str(status) in REVIEW_STATUSES for status in statuses)
            values["review_count"] = review_count
            values["review_rate"] = round(review_count / len(statuses) * 100, 1)
        return values

    @staticmethod
    def _total_tokens(usage: dict[str, Any]) -> int | None:
        total_tokens = usage.get("total_tokens")
        if isinstance(total_tokens, int) and not isinstance(total_tokens, bool):
            return total_tokens
        token_values = [
            value
            for value in (usage.get("input_tokens"), usage.get("output_tokens"))
            if isinstance(value, int) and not isinstance(value, bool)
        ]
        return sum(token_values) if token_values else None

    def _serialize(
        self,
        value: dict[str, Any],
        include_items: bool = False,
    ) -> dict[str, Any]:
        source = json_value(value.pop("source_json"), {})
        snapshot = json_value(value.pop("snapshot_json", None), {})
        value.pop("sample_json", None)
        items = json_value(value.pop("result_json"), [])
        value["summary"] = json_value(value.pop("summary_json"), {})
        value["usage"] = json_value(value.pop("usage_json"), {})
        value["metrics"] = json_value(value.pop("metrics_json"), {})
        if "efficiency" not in value["summary"]:
            efficiency = self._efficiency_summary(
                items,
                value["usage"],
                value["metrics"],
            )
            if efficiency["sides"]:
                value["summary"]["efficiency"] = efficiency
        value["model_names"] = json_value(value.pop("model_names_json"), [])
        value["source"] = source.get("result", {})
        if include_items:
            value["items"] = items
        with self.database.connect() as connection:
            draft = connection.execute(
                """
                SELECT revision FROM classification_standard_drafts WHERE id = ?
                """,
                (value["draft_id"],),
            ).fetchone()
            approver = (
                connection.execute(
                    "SELECT display_name FROM users WHERE id = ?",
                    (value.get("approved_by"),),
                ).fetchone()
                if value.get("approved_by")
                else None
            )
        value["approved_by_name"] = (
            str(approver["display_name"]) if approver is not None else None
        )
        value["is_current"] = bool(
            draft
            and int(draft["revision"]) == int(value["draft_revision"])
            and (
                source.get("comparison_type", "standard_version") != "standard_version"
                or validation_contract_matches(snapshot, source)
            )
        )
        if items:
            value["summary"]["reference_evaluation"] = evaluate_references(items)
            if source.get("comparison_type", "standard_version") == "standard_version":
                value["summary"]["reference_evaluation"]["sides"].pop("baseline", None)
        value["quality_gate"] = publication_quality_gate(
            snapshot, source, [], value["summary"]
        )
        value["publication_ready"] = bool(
            value["quality_gate"]["passed"]
            and value["is_current"]
            and value["status"] == "completed"
            and int(value["error_count"]) == 0
            and bool(value.get("approved_at"))
            and source.get("comparison_type", "standard_version") == "standard_version"
        )
        return value
