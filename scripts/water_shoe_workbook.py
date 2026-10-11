from __future__ import annotations

from pathlib import Path
from typing import cast

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from scripts.water_shoe_workbook_layout import (
    CLAIM_HEADER_COMMENTS,
    CLAIM_HEADERS,
    MISSING_FILL,
    SKU_HEADER_COMMENTS,
    SKU_HEADERS,
    add_list_validation,
    style_sheet,
)


def build_workbook(rows: list[dict[str, object]], output_path: Path) -> None:
    workbook = Workbook()
    sku_sheet = cast(Worksheet, workbook.active)
    sku_sheet.title = "SKU客观特征"
    style_sheet(
        sku_sheet,
        SKU_HEADERS,
        SKU_HEADER_COMMENTS,
        prefilled_count=8,
        required_count=7,
    )

    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(SKU_HEADERS, start=1):
            sku_sheet.cell(row=row_index, column=column_index).value = cast(
                str | int | None, row.get(header)
            )
        sku_sheet.cell(row=row_index, column=15).number_format = "yyyy-mm-dd"

    last_sku_row = len(rows) + 1
    sku_table = Table(displayName="SkuObjectiveFacts", ref=f"A1:X{last_sku_row}")
    sku_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sku_sheet.add_table(sku_table)

    add_list_validation(
        sku_sheet,
        f"K2:K{last_sku_row}",
        ["US_UNISEX", "US_WOMEN", "US_MEN", "US_KIDS", "EU", "UK", "OTHER"],
    )
    add_list_validation(
        sku_sheet,
        f"N2:N{last_sku_row}",
        ["product_spec", "sample_measurement", "supplier_spec", "lab_test", "mixed"],
    )
    add_list_validation(
        sku_sheet,
        f"P2:Q{last_sku_row}",
        ["narrow", "standard", "wide", "unknown"],
    )
    sku_sheet.conditional_formatting.add(
        f"I2:O{last_sku_row}",
        FormulaRule(formula=['I2=""'], fill=MISSING_FILL),
    )

    claim_sheet = workbook.create_sheet("Listing声明")
    style_sheet(
        claim_sheet,
        CLAIM_HEADERS,
        CLAIM_HEADER_COMMENTS,
        prefilled_count=0,
        required_count=8,
    )
    claim_sheet.append([None] * len(CLAIM_HEADERS))
    for column_index in (6, 7, 8):
        claim_sheet.cell(row=2, column=column_index).number_format = "yyyy-mm-dd"

    claim_table = Table(displayName="ListingClaims", ref="A1:I2")
    claim_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    claim_sheet.add_table(claim_table)
    add_list_validation(
        claim_sheet,
        "C2:C10000",
        [
            "size_fit",
            "width_fit",
            "material",
            "quick_dry",
            "drainage",
            "slip_resistance",
            "breathability",
            "protection",
            "intended_use",
            "package_quantity",
            "other",
        ],
    )
    add_list_validation(
        claim_sheet,
        "E2:E10000",
        ["title", "bullet", "A+", "size_chart", "package", "manual", "other"],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def verify_workbook(output_path: Path, expected_rows: int) -> None:
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    if workbook.sheetnames != ["SKU客观特征", "Listing声明"]:
        raise ValueError(f"工作表不符合预期: {workbook.sheetnames}")

    sku_sheet = workbook["SKU客观特征"]
    if sku_sheet.max_row != expected_rows + 1:
        raise ValueError("SKU 模板行数与生成结果不一致")

    headers = [cell.value for cell in sku_sheet[1]]
    if headers != SKU_HEADERS:
        raise ValueError("SKU 模板字段与契约不一致")

    sku_values = [
        sku_sheet.cell(row=row, column=1).value
        for row in range(2, sku_sheet.max_row + 1)
    ]
    if len(sku_values) != len(set(sku_values)):
        raise ValueError("SKU 模板中存在重复 canonical_sku")
