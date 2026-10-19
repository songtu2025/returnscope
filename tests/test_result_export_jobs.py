from __future__ import annotations

import json
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event

import pytest
from openpyxl import load_workbook
from test_classification_result_download import context as context

from web_backend.classification_results.export_jobs import ResultExportService


@pytest.fixture
def exports(context, tmp_path):
    service = ResultExportService(context.service, tmp_path / "exports")
    yield service
    service.close()


def test_concurrent_requests_reuse_one_job_and_completed_file(
    context, exports, monkeypatch
):
    version_ids = [context.version["version_id"]]
    captured = Barrier(3)
    original = exports.results.capture_export

    def capture(ids):
        snapshot = original(ids)
        captured.wait(timeout=3)
        return snapshot

    monkeypatch.setattr(exports.results, "capture_export", capture)
    with ThreadPoolExecutor(max_workers=3) as executor:
        jobs = list(executor.map(exports.create, [version_ids] * 3))
    monkeypatch.setattr(exports.results, "capture_export", original)
    assert len({item["id"] for item in jobs}) == 1
    assert all(item["status"] == "queued" for item in jobs)
    assert len(list(exports.directory.glob("*.json"))) == 1
    assert len(list(exports.directory.glob("*.snapshot.gz"))) == 1
    assert not list(exports.directory.glob("*.tmp"))
    job_id = jobs[0]["id"]
    assert not exports._path(job_id, ".xlsx").exists()
    assert exports.claim_next() == job_id
    assert exports.claim_next() is None
    exports.run(job_id)
    path, filename = exports.file(job_id)
    assert filename == "classification-L1-v1.xlsx"
    assert exports.create(version_ids)["status"] == "ready"
    book = load_workbook(path)
    assert book.sheetnames == ["反馈汇总", "语义明细", "源明细"]
    assert book["反馈汇总"].max_row == 3
    assert book["源明细"].max_row == 4
    book.close()


@pytest.mark.parametrize("stage", ["capture", "compress"])
def test_preparing_snapshot_does_not_block_existing_job_queries(
    context, exports, monkeypatch, stage
):
    from web_backend.classification_results import export_jobs

    versions = [context.version["version_id"]]
    existing = exports.create(versions)
    with context.database.transaction() as connection:
        connection.execute(
            "UPDATE classification_result_records SET comment='合成更新内容'"
        )
    entered, release = Event(), Event()
    target, name = (
        (exports.results, "capture_export")
        if stage == "capture"
        else (export_jobs.gzip, "compress")
    )
    original = getattr(target, name)

    def prepare(value):
        entered.set()
        if not release.wait(5):
            raise RuntimeError("合成快照释放超时")
        return original(value)

    monkeypatch.setattr(target, name, prepare)
    with ThreadPoolExecutor(max_workers=2) as executor:
        creating = executor.submit(exports.create, versions)
        try:
            assert entered.wait(3)
            querying = executor.submit(exports.get, existing["id"])
            assert querying.result(timeout=1)["id"] == existing["id"]
        finally:
            release.set()
        created = creating.result(timeout=3)
        assert created["id"] != existing["id"]


def test_snapshot_survives_edits_and_new_click_exports_latest(context, exports):
    version_ids = [context.version["version_id"]]
    old = exports.create(version_ids)
    with context.database.transaction() as connection:
        connection.execute(
            "UPDATE classification_result_records SET comment='修改后的反馈'"
        )
    new = exports.create(version_ids)
    assert old["id"] != new["id"]
    for job, expected in [(old, "Too small"), (new, "修改后的反馈")]:
        exports.run(job["id"])
        path, _ = exports.file(job["id"])
        book = load_workbook(path)
        for sheet_name in ["反馈汇总", "源明细"]:
            values = list(book[sheet_name].values)
            index = values[0].index("反馈正文")
            assert all(row[index] == expected for row in values[1:])
        book.close()


def test_failed_generation_can_retry_and_running_job_recovers(
    context, exports, monkeypatch
):
    versions = [context.version["version_id"]]
    job = exports.create(versions)
    original = exports.results.snapshot_workbook

    def fail(_snapshot):
        raise ValueError("不应泄露的原文")

    monkeypatch.setattr(exports.results, "snapshot_workbook", fail)
    exports.run(job["id"])
    failed = exports.get(job["id"])
    assert failed["status"] == "failed"
    assert "不应泄露" not in json.dumps(failed, ensure_ascii=False)
    assert exports.create(versions)["status"] == "queued"
    assert exports.claim_next() == job["id"]
    restarted = ResultExportService(context.service, exports.directory)
    restarted.recover()
    assert restarted.get(job["id"])["status"] == "queued"
    monkeypatch.setattr(exports.results, "snapshot_workbook", original)
    restarted.run(job["id"])
    assert restarted.get(job["id"])["status"] == "ready"


def test_expiration_cleans_only_export_files(context, exports):
    job = exports.create([context.version["version_id"]])
    exports.run(job["id"])
    unrelated = exports.directory / "unrelated.txt"
    unrelated.write_text("保留", encoding="utf-8")
    manifest = exports._load(job["id"])
    manifest["expires_at"] = 0
    exports._save(manifest)
    with pytest.raises(FileNotFoundError, match="过期"):
        exports.file(job["id"])
    exports.claim_next()
    assert unrelated.exists()
    assert not exports._path(job["id"]).exists()
    assert not exports._path(job["id"], ".xlsx").exists()
    assert not exports._path(job["id"], ".snapshot.gz").exists()


def test_permanent_result_deletion_blocks_cached_download(context, exports):
    job = exports.create([context.version["version_id"]])
    exports.run(job["id"])
    with context.database.transaction() as connection:
        connection.execute("DELETE FROM classification_result_records")
        connection.execute("DELETE FROM classification_units")
        connection.execute("UPDATE task_segments SET result_version_id=NULL")
        connection.execute("DELETE FROM classification_result_versions")
    from web_backend.classification_result_service import ClassificationResultNotFound

    with pytest.raises(ClassificationResultNotFound):
        exports.file(job["id"])


@pytest.mark.parametrize("generation_fails", [False, True])
def test_generation_crossing_expiration_finishes_and_can_retry(
    context, exports, monkeypatch, generation_fails
):
    from web_backend.classification_results import export_jobs

    clock = [2_000_000_000.0]
    monkeypatch.setattr(export_jobs.time, "time", lambda: clock[0])
    versions = [context.version["version_id"]]
    job = exports.create(versions)
    assert exports.claim_next() == job["id"]
    original = exports.results.snapshot_workbook

    def generate(snapshot):
        clock[0] = job["expires_at"] + 1
        if generation_fails:
            raise ValueError("合成生成失败")
        return original(snapshot)

    monkeypatch.setattr(exports.results, "snapshot_workbook", generate)
    exports.run(job["id"])
    with pytest.raises(FileNotFoundError, match="过期"):
        exports.file(job["id"])
    for suffix in (".json", ".snapshot.gz", ".xlsx", ".xlsx.tmp"):
        assert not exports._path(job["id"], suffix).exists()
    monkeypatch.setattr(exports.results, "snapshot_workbook", original)
    retry = exports.create(versions)
    assert retry["status"] == "queued"
    assert retry["expires_at"] > clock[0]
    assert exports.claim_next() == retry["id"]
    exports.run(retry["id"])
    assert exports.get(retry["id"])["status"] == "ready"


@pytest.mark.parametrize("status", ["queued", "running", "ready"])
def test_existing_export_skips_compression(context, exports, monkeypatch, status):
    from unittest.mock import Mock

    from web_backend.classification_results import export_jobs

    versions = [context.version["version_id"]]
    existing = exports.create(versions)
    if status in {"running", "ready"}:
        assert exports.claim_next() == existing["id"]
    if status == "ready":
        exports.run(existing["id"])
    compress = Mock(wraps=export_jobs.gzip.compress)
    monkeypatch.setattr(export_jobs.gzip, "compress", compress)
    repeated = exports.create(versions)
    assert repeated["id"] == existing["id"]
    assert repeated["status"] == status
    compress.assert_not_called()
